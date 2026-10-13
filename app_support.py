"""Safe local exports and diagnostics shared by the desktop application."""

from contextlib import contextmanager
from datetime import datetime
import logging
from logging.handlers import RotatingFileHandler
import os
import tempfile

import pandas as pd


@contextmanager
def atomic_output(destination):
    """Replace a destination only after its complete export succeeds."""
    destination = os.path.abspath(os.fspath(destination))
    directory = os.path.dirname(destination)
    suffix = os.path.splitext(destination)[1]
    descriptor, temporary = tempfile.mkstemp(prefix='.signal-export-', suffix=suffix, dir=directory)
    os.close(descriptor)
    try:
        yield temporary
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)


def write_excel_frame(writer, frame, sheet_name):
    """Write every row, splitting large datasets over numbered worksheets."""
    if len(frame.columns) > 16_384:
        raise ValueError('This dataset has more columns than an Excel worksheet supports. Use CSV export.')
    export = frame.copy(deep=False)
    timezone_columns = [column for column in frame.columns
                        if isinstance(frame[column].dtype, pd.DatetimeTZDtype)]
    if timezone_columns:
        export = frame.copy()
        for column in timezone_columns:
            export[column] = export[column].dt.tz_convert('UTC').dt.tz_localize(None)
    maximum_rows = 1_048_575  # Leave one row for the column headings.
    parts = max(1, (len(export) + maximum_rows - 1) // maximum_rows)
    names = []
    for index in range(parts):
        suffix = f'_{index + 1}' if parts > 1 else ''
        name = sheet_name[:31 - len(suffix)] + suffix
        duplicate = 1
        while name in writer.sheets:
            collision_suffix = f'_{duplicate}'
            name = sheet_name[:31 - len(collision_suffix)] + collision_suffix
            duplicate += 1
        export.iloc[index * maximum_rows:(index + 1) * maximum_rows].to_excel(
            writer, sheet_name=name, index=False
        )
        writer.sheets[name].freeze_panes = 'A2'
        writer.sheets[name].auto_filter.ref = writer.sheets[name].dimensions
        if sheet_name == 'Scorecard_Summary' and 'status' in export:
            from openpyxl.styles import Font, PatternFill
            colours = {'GREEN':('E8F5E9','007E33'), 'YELLOW':('FFF3B0','725300'),
                       'RED':('FDE8E7','B42318')}
            sheet = writer.sheets[name]
            section = export.iloc[index * maximum_rows:(index + 1) * maximum_rows]
            for row_index, status in enumerate(section['status'], start=2):
                if status in colours:
                    fill, foreground = colours[status]
                    for cell in sheet[row_index]:
                        cell.fill = PatternFill(fill_type='solid', fgColor=fill)
                        cell.font = Font(color=foreground)
            for position, column in enumerate(export.columns, start=1):
                if column == 'status':
                    sheet.cell(1, position).value = 'colour_band'
            for letter, width in {'A':38,'B':24,'C':24,'D':22,'E':18,'F':78,'G':18}.items():
                sheet.column_dimensions[letter].width = width
        names.append(name)
    return names


def configure_logging(directory):
    logger = logging.getLogger('appliance_signal_studio')
    if not logger.handlers:
        try:
            handler = RotatingFileHandler(
                os.path.join(directory, 'application.log'), encoding='utf-8',
                maxBytes=2_000_000, backupCount=2
            )
            handler.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(message)s'))
            logger.addHandler(handler)
        except OSError:
            logger.addHandler(logging.NullHandler())
        logger.setLevel(logging.INFO)
        logger.propagate = False
    return logger


def export_timestamp():
    return datetime.now().astimezone().isoformat(timespec='seconds')
