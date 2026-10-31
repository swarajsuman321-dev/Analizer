"""Readable comparison reports with complete, separately inspectable evidence.

The PDF is a narrative of the current workspace or an explicitly selected pair.
Raw interval samples, all difference intervals and all events remain in the
workbook/JSON. A report never infers pairings, calibrations or acceptance limits.
"""
from copy import copy, deepcopy
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
import math
from pathlib import Path
import textwrap
from types import SimpleNamespace
from uuid import uuid4

import numpy as np
import pandas as pd

from app_operation_context import checkpoint
from app_presentation import CURSOR_COLORS, match_color
from app_channel_identity import channel_label
from app_version import VERSION, BUILD_DATE


NOTES = ('Prepared by', 'Reviewed by', 'Test conditions', 'Observations',
         'Conclusions', 'Recommended actions')
INK = '#182c42'
MUTED = '#526074'
EXCEL_DATA_ROWS = 1048575


def clean(value):
    """Produce strict JSON without silently turning missing data into zero."""
    if isinstance(value, dict):
        return {str(key): clean(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [clean(item) for item in value]
    if isinstance(value, (float, np.floating)):
        return float(value) if np.isfinite(value) else None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer)):
        return int(value)
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, (str,)):
        return value
    return str(value)


def text(value):
    if value is None or value is pd.NA or value is pd.NaT:
        return 'Not Available'
    if isinstance(value, (float, np.floating)):
        return f'{value:.6g}' if np.isfinite(value) else 'Not Available'
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(clean(value), ensure_ascii=False, sort_keys=True)
    return str(value)


def _snapshot_hash(channel):
    digest = hashlib.sha256()
    for name in ('x', 'y', 'raw_y'):
        values = getattr(channel, name)
        digest.update(name.encode())
        if values is not None:
            digest.update(memoryview(np.ascontiguousarray(values)).cast('B'))
    digest.update(json.dumps(clean({'unit': channel.unit, 'kind': channel.kind,
        'coordinate_kind': channel.coordinate_kind, 'coordinate_unit': channel.coordinate_unit,
        'shift': channel.shift, 'masks': channel.masks, 'parents': channel.parents,
        'recipe': channel.recipe}), sort_keys=True).encode())
    return digest.hexdigest()


def visible_views(workspace):
    from app_plot_workspace import LAYOUTS
    rows, columns = LAYOUTS[workspace['layout']]
    return list(workspace['panes'][:rows*columns])


def _source_rows(store, ids):
    rows = []
    for dataset_id in dict.fromkeys(store.channels[id].dataset_id for id in ids):
        members = [store.channels[id] for id in ids if store.channels[id].dataset_id == dataset_id]
        dataset = store.datasets.get(dataset_id)
        source = dataset.source if dataset else members[0].source
        record = {'dataset_id': dataset_id, 'dataset': dataset.name if dataset else 'Unassigned dataset',
                  'source': source or 'Source path not retained', 'selected_channels': len(members),
                  'source_sha256': None, 'fingerprint_method': 'Source bytes unavailable; see loaded channel snapshots',
                  'source_bytes': None}
        try:
            path = Path(source)
            if source and path.is_file():
                digest = hashlib.sha256()
                with path.open('rb') as stream:
                    while block := stream.read(1024*1024):
                        checkpoint('Fingerprinting '+path.name)
                        digest.update(block)
                record.update(source_sha256=digest.hexdigest(), source_bytes=path.stat().st_size,
                    fingerprint_method='Current source file bytes at report generation; not an import-time identity claim')
        except (OSError, ValueError):
            record['fingerprint_method'] = 'Source file unreadable; see loaded channel snapshots'
        rows.append(record)
    return rows


def _quality_rows(store, workspace):
    records = []
    for pane, view in enumerate(visible_views(workspace), 1):
        a, b = sorted(view.get('cursors', [0, 0]))
        for id in view['ids']:
            checkpoint('Inspecting channel timing quality')
            channel = store.channels[id]
            x, y = channel.values()
            row = {'pane': pane, 'channel_id': id, 'channel': channel_label(store, channel),
                   'unit': channel.unit, 'coordinate_kind': channel.coordinate_kind,
                   'coordinate_unit': channel.coordinate_unit,
                   'requested_start': a, 'requested_end': b,
                   'mask_count': len(channel.masks), 'source_samples': len(x)}
            if channel.coordinate_kind == 'time':
                from app_engineering_trace import Trace
                from app_time_units import minute_factor
                factor = minute_factor(channel.coordinate_unit)*60
                trace = Trace(channel)
                stats = trace.stats(a*factor, b*factor)
                row.update(trace.quality(a*factor, b*factor))
                row.update(supported_mean=stats.get('mean'), supported_rms=stats.get('rms'))
            else:
                selected = (x >= a) & (x <= b)
                row.update(recorded_samples=int(selected.sum()),
                    missing_readings=int((~np.isfinite(y[selected])).sum()),
                    duplicate_timestamps=int(len(x[selected])-len(np.unique(x[selected]))),
                    interval_policy='Native non-time coordinate; elapsed-time quality is not applicable',
                    quality_flags=['Elapsed-time coverage not applicable'])
            records.append(clean(row))
    return records


def _flatten(value, prefix=''):
    rows = []
    if isinstance(value, dict):
        for key, item in value.items():
            rows.extend(_flatten(item, prefix+'.'+str(key) if prefix else str(key)))
    else:
        value=clean(value)
        rows.append({'Field': prefix, 'Value': text(value) if isinstance(value,(dict,list)) else value})
    return rows


def worksheet_parts(name, count, limit=None):
    """Plan complete, header-repeated worksheets without writing giant fixtures."""
    limit=EXCEL_DATA_ROWS if limit is None else limit
    if limit<1:raise ValueError('Worksheet data-row limit must be positive')
    return [(name if index==0 else name[:31-len(' '+str(index+1))]+' '+str(index+1),
             start,min(count,start+limit))
            for index,start in enumerate(range(0,max(1,count),limit))]


def build_evidence(store, workspace, data, summary, metadata=None, scorecard=None,
                   title='Signal comparison report', report_id=None):
    """Freeze shared tables once so PDF, workbook and JSON agree exactly."""
    metadata = deepcopy(metadata or {})
    ids = list(dict.fromkeys(id for view in visible_views(workspace) for id in view['ids']))
    pair = clean(metadata.get('Signal comparison'))
    info = {'Report ID': report_id or uuid4().hex, 'Report title': metadata.get('Report title') or title,
            'Application version': VERSION, 'Build date': BUILD_DATE,
            'Generated UTC': datetime.now(timezone.utc).isoformat(),
            'Workspace': workspace.get('workspace_name', 'Comparison'),
            'Statistics scope': workspace.get('statistics_scope', 'Cursor interval'),
            'Pairing policy': 'Explicit reference/candidate selection' if pair else
                'No pair selected; channels are source-qualified independent measurements',
            'Data policy': 'Original interval samples retained. No resampling or extrapolation in raw exports.',
            'Cursor statistics policy': 'Workspace interpolation/snap and locked original-sample identity; quality timing policy is reported separately.',
            'Metadata': clean({k: v for k, v in metadata.items() if k != 'Signal comparison'})}
    sources = _source_rows(store, ids)
    quality = _quality_rows(store, workspace)
    colors = _colors(store)
    provenance = [{'channel_id': id, 'channel': channel_label(store, id),
        'dataset_id': store.channels[id].dataset_id,
        'source_column': store.channels[id].attributes.get('source_column', store.channels[id].name),
        'unit': store.channels[id].unit,
        'coordinate_kind': store.channels[id].coordinate_kind,
        'coordinate_unit': store.channels[id].coordinate_unit,
        'original_coordinate_unit': store.channels[id].attributes.get('original_coordinate_unit'),
        'shift': store.channels[id].shift, 'masks': store.channels[id].masks,
        'parents': store.channels[id].parents, 'recipe': store.channels[id].recipe,
        'trace_color': colors[id],
        'plot_style': deepcopy(store.channels[id].attributes.get('plot_style', {})),
        'loaded_snapshot_sha256': _snapshot_hash(store.channels[id]),
        'fingerprint_method': 'Loaded raw coordinate/value arrays, unit/clock, shift, masks and processing recipe; appearance excluded'} for id in ids]
    views = [{'pane': index, **clean(view)} for index, view in enumerate(visible_views(workspace), 1)]
    evidence = clean({'schema': 1, 'report': info, 'sources': sources, 'workspace': workspace,
        'channel_quality': quality, 'processing_provenance': provenance,
        'cursor_summary': summary.to_dict('records'), 'signal_comparison': pair,
        'scorecard': scorecard.to_dict('records') if scorecard is not None else [],
        'raw_interval_sample_count': len(data),
        'raw_interval_samples': data.to_dict('records')})
    tables = {'Report details': pd.DataFrame(_flatten(info)), 'Sources': pd.DataFrame(sources),
        'Channel quality': pd.DataFrame(quality), 'View settings': pd.DataFrame(views),
        'Processing provenance': pd.DataFrame(provenance)}
    if pair:
        tables['Pair comparison'] = pd.DataFrame(_flatten({key: value for key, value in pair.items()
            if key not in ('differences', 'events', 'reference', 'candidate')}))
        tables['Pair source statistics'] = pd.DataFrame([
            {'Role': role, **row} for role in ('reference', 'candidate')
            for row in _flatten({key: value for key, value in pair.get(role, {}).items() if key != 'plot'})])
        tables['Difference intervals'] = pd.DataFrame(pair.get('differences') or [])
        tables['Activity events'] = pd.DataFrame(pair.get('events') or [])
    return evidence, tables


def _excel_value(value):
    """Match the existing pandas export types without allocating Excel cells.

    Missing readings remain explicit text. Source strings, including formulas
    and error tokens, are written literally by the streaming writer below.
    """
    # DataFrame.itertuples yields native Python primitives for ordinary
    # measurement columns. Keep their common path cheap across millions of
    # cells; extension/object columns fall through to the explicit conversions.
    kind = type(value)
    if kind is str:return value, None
    if kind is float:
        if math.isnan(value):return 'Not Available', None
        if math.isinf(value):return '-inf' if value < 0 else 'inf', None
        return value, None
    if kind is int or kind is bool:return value, None
    if isinstance(value, (dict, list, tuple, np.ndarray)):
        return text(value), None
    if value is None or value is pd.NA or value is pd.NaT:
        return 'Not Available', None
    if isinstance(value, (bool, np.bool_)):
        return bool(value), None
    if isinstance(value, (int, np.integer)):
        return int(value), None
    if isinstance(value, (float, np.floating)):
        if np.isnan(value):return 'Not Available', None
        if np.isinf(value):return '-inf' if value < 0 else 'inf', None
        return float(value), None
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            raise ValueError('Excel does not support datetimes with timezones.')
        return value, 'YYYY-MM-DD HH:MM:SS'
    if isinstance(value, date):
        return value, 'YYYY-MM-DD'
    if isinstance(value, timedelta):
        return value.total_seconds()/86400, '0'
    return str(value), None


def write_workbook(path, data, summary, evidence, tables, scorecard=None):
    """Stream complete evidence rows with bounded Excel-cell memory.

    A realistic comparison has millions of cells. Building a normal worksheet
    and revisiting every cell retained them all in memory. Append each row
    once to bound Excel-cell memory, preserving the
    same sheets, full numeric precision, literal text and match accents.
    """
    from contextlib import suppress
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    frames = {'Active interval': data, 'Cursor summary': summary}
    if scorecard is not None:
        frames['Comparison'] = scorecard
    frames.update(tables)
    parts = [{'Table':name, 'Sheet':sheet, 'First source row':start+1,
              'Last source row':end, 'Data rows':end-start}
             for name,frame in frames.items()
             for sheet,start,end in worksheet_parts(name,len(frame))]
    evidence['workbook_parts'] = parts
    frames['Workbook parts'] = pd.DataFrame(parts)
    extended = []
    book = Workbook(write_only=True)
    heading_font = Font(bold=True, color='FFFFFF')
    heading_fill = PatternFill('solid', fgColor='243A53')
    match_fills = {}

    def write_part(name, frame, start, end, *, retain_extended=True):
        sheet = book.create_sheet(name)
        columns = list(frame.columns)
        # Write-only XML normally omits dimensions. The exact size is known
        # before streaming: publish it so read-only Excel consumers can inspect
        # row counts without scanning or loading the complete worksheet.
        dimension = 'A1:'+get_column_letter(max(1,len(columns)))+str(end-start+1)
        sheet.calculate_dimension = lambda value=dimension: value
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = dimension
        for number,column in enumerate(columns,1):
            values = [column, *frame.iloc[start:min(end,start+60),number-1].tolist()]
            sheet.column_dimensions[get_column_letter(number)].width = min(60,max(13,
                max(len(str(value)) for value in values)+2))
        headings = []
        for column in columns:
            cell = WriteOnlyCell(sheet, str(column))
            cell.data_type = 's'
            cell.font, cell.fill = heading_font, heading_fill
            headings.append(cell)
        sheet.append(headings)
        match_index = columns.index('match_percent') if 'match_percent' in columns else -1
        for position, values in enumerate(frame.iloc[start:end].itertuples(index=False,name=None),start):
            if (position-start)%4096 == 0:
                checkpoint('Writing '+name, position-start, end-start)
            cells = []
            for number,value in enumerate(values):
                value, number_format = _excel_value(value)
                if retain_extended and isinstance(value,str) and len(value)>30000:
                    identity = uuid4().hex
                    for part,offset in enumerate(range(0,len(value),30000),1):
                        extended.append({'Text ID':identity, 'Sheet':name,
                            'Source row':position+2, 'Column':columns[number],
                            'Part':part, 'Text':value[offset:offset+30000]})
                    value = 'Complete text: Extended text / '+identity
                color = match_color(value) if number==match_index else None
                if isinstance(value,str) or number_format or color:
                    cell = WriteOnlyCell(sheet,value)
                    if isinstance(value,str):cell.data_type='s'
                    if number_format:cell.number_format=number_format
                    if color:
                        if color not in match_fills:
                            match_fills[color] = PatternFill('solid',fgColor=color.lstrip('#'))
                        cell.fill = match_fills[color]
                        cell.number_format = '0.0"%"'
                    cells.append(cell)
                else:
                    cells.append(value)
            sheet.append(cells)
        sheet.close()

    try:
        for original_name, frame in frames.items():
            for name,start,end in worksheet_parts(original_name,len(frame)):
                checkpoint('Writing '+name,0,end-start)
                write_part(name,frame,start,end)
        if extended:
            frame = pd.DataFrame(extended)
            for name,start,end in worksheet_parts('Extended text',len(frame)):
                write_part(name,frame,start,end,retain_extended=False)
        checkpoint('Finalizing complete comparison workbook')
        book.save(path)
    finally:
        # Cancellation before save must not leave open write-only XML streams
        # or their temporary files. The destination itself is staged atomically.
        for sheet in book:
            with suppress(Exception):
                if not sheet.closed:sheet.close()
            writer = getattr(sheet,'_writer',None)
            if writer is not None:
                with suppress(OSError):
                    writer.cleanup()
        book.close()


class _Pages:
    """Small deterministic portrait-page layout; text never relies on auto-fit."""
    def __init__(self, pdf, report):
        self.pdf, self.report = pdf, report
        self.figure = None
        self.page_number = 0

    def page(self, section, subtitle=''):
        self.save()
        from matplotlib.figure import Figure
        self.figure = Figure(figsize=(8.27, 11.69))
        self.page_number += 1
        self.section = section
        self.y = .89
        heading=textwrap.shorten(str(self.report['Report title']).replace('\n',' '),width=80,placeholder='…')
        self.figure.text(.07, .963, heading, color=INK, fontsize=15,
                         weight='bold', va='top', parse_math=False)
        self.figure.text(.07, .925, section+(' · '+subtitle if subtitle else ''), color=MUTED,
                         fontsize=10, va='top', parse_math=False)
        self.figure.text(.07, .025, f"Report {self.report['Report ID']} · {self.report['Generated UTC']} · Page {self.page_number}",
                         fontsize=6.5, color=MUTED, parse_math=False)

    def ensure(self, height):
        if self.y-height < .07:
            self.page(self.section, 'continued')

    def paragraph(self, value, *, size=9, label=None):
        if label:
            self.ensure(.035)
            self.figure.text(.07, self.y, label, fontsize=10, weight='bold', color=INK,
                             va='top', parse_math=False)
            self.y -= .025
        for original in str(value).splitlines() or ['']:
            lines = textwrap.wrap(original, width=max(40, int(112*9/size)),
                                  replace_whitespace=False, drop_whitespace=False) or ['']
            for line in lines:
                self.ensure(.018)
                self.figure.text(.07, self.y, line, fontsize=size, color=INK, va='top', parse_math=False)
                self.y -= .018*(size/9)
        self.y -= .012

    def table(self, rows, columns, widths=None):
        """Repeat headers and split long cells across readable rows, without truncation."""
        widths = widths or [1/len(columns)]*len(columns)
        def header():
            self.ensure(.05)
            x=.07
            for name, width in zip(columns, widths):
                self.figure.text(x, self.y, name, fontsize=8, weight='bold', color=INK,
                                 va='top', parse_math=False)
                x += .86*width
            self.y -= .033
        header()
        for record in rows:
            checkpoint('Paginating comparison evidence')
            cells = [textwrap.wrap(text(record.get(column)), max(8, int(112*width)),
                     replace_whitespace=False, drop_whitespace=False) or ['']
                     for column, width in zip(columns, widths)]
            for offset in range(0, max(map(len, cells)), 8):
                count = max(len(cell[offset:offset+8]) for cell in cells)
                height = .015*count+.013
                if self.y-height < .07:
                    self.page(self.section, 'continued')
                    header()
                x=.07
                for cell, width in zip(cells, widths):
                    self.figure.text(x, self.y, '\n'.join(cell[offset:offset+8]),
                        fontsize=7.7, color=INK, va='top', linespacing=1.2, parse_math=False)
                    x += .86*width
                self.y -= height
        self.y -= .015

    def save(self):
        if self.figure is not None:
            checkpoint('Writing comparison PDF page '+str(self.page_number))
            self.pdf.savefig(self.figure,dpi=180)
            self.figure.clear()
            self.figure = None


def _colors(store):
    """Use the UI palette without creating appearance metadata during export."""
    from app_trace_colors import channel_colors
    channels = {}
    for id, channel in store.channels.items():
        value = copy(channel)
        value.attributes = deepcopy(channel.attributes)
        channels[id] = value
    return channel_colors(SimpleNamespace(channels=channels))


def _metric_rows(pair):
    metrics=pair.get('metrics',{})
    unit=metrics.get('unit') or ''
    names={'unit':'Comparison value unit','compared_duration_min':'Supported comparison time (min)',
        'coverage_fraction':'Supported comparison coverage (%)',
        'mean_bias':'Time-weighted average difference ('+unit+')',
        'mean_absolute_difference':'Mean absolute difference ('+unit+')',
        'rmse':'Root mean square difference ('+unit+')',
        'max_absolute_difference':'Largest absolute difference ('+unit+')',
        'max_difference_time_min':'Time of largest difference (elapsed min)',
        'different_duration_min':'Supported time with distinguishable differences (min)',
        'tolerance':'Configured comparison tolerance ('+unit+')',
        'within_tolerance_duration_min':'Supported time within configured tolerance (min)'}
    # The first-page table is an executive summary. Full definitions, floating
    # resolution and integration policies stay in the complete workbook/JSON;
    # adding an engine evidence field must not expand this table by accident.
    keys=['unit','compared_duration_min','coverage_fraction','mean_bias',
          'mean_absolute_difference','rmse','max_absolute_difference',
          'max_difference_time_min','different_duration_min']
    if metrics.get('tolerance') is not None:
        keys.extend(('tolerance','within_tolerance_duration_min'))
    return [{'Measure':names[key],
             'Value':text(100*metrics[key] if key=='coverage_fraction' and metrics[key] is not None else metrics[key])}
            for key in keys if key in metrics]


def _paired_plot(pages, pair,store,colors):
    pages.page('2. Graphs', 'Elapsed comparison of the explicitly selected signals')
    axes = pages.figure.subplots(2, 1)
    pages.figure.subplots_adjust(left=.12, right=.93, bottom=.14, top=.85, hspace=.32)
    for role, fallback in zip(('reference', 'candidate'), ('#075985', '#b45309')):
        source=pair.get(role, {})
        color=colors.get(source.get('channel_id'),fallback)
        plot=source.get('plot') or {}
        channel=store.channels.get(source.get('channel_id'))
        binary=bool(channel and channel.attributes.get('actuator_binary'))
        axes[0].plot(plot.get('time_min', []), plot.get('value', []), color=color,
            label=role.title()+': '+str(source.get('dataset', source.get('name', 'Source'))),
            drawstyle='steps-post' if binary else 'default', linewidth=.9)
    unit=pair.get('metrics', {}).get('unit') or pair.get('reference', {}).get('unit', '')
    axes[0].set_ylabel(unit or 'Value',parse_math=False)
    if all(store.channels.get(pair.get(role,{}).get('channel_id')) is not None and
           store.channels[pair[role]['channel_id']].attributes.get('actuator_binary')
           for role in ('reference','candidate')):
        axes[0].set_yticks([0,1],['OFF','ON'])
    legend=axes[0].legend(fontsize=7)
    for label in legend.get_texts():label.set_parse_math(False)
    times=[];values=[]
    for row in pair.get('differences') or []:
        if row.get('supported'):
            times.extend([row.get('start_min'), row.get('end_min'), np.nan])
            values.extend([row.get('delta_start'), row.get('delta_end'), np.nan])
    axes[1].plot(times, values, color='#7e22ce', linewidth=.9,rasterized=True)
    axes[1].axhline(0, color=MUTED, linewidth=.5)
    axes[1].set_ylabel('Candidate − Reference\n'+str(unit),parse_math=False)
    for axis in axes:
        axis.set_xlabel('Elapsed time (min)')
        axis.grid(True, alpha=.18)
        axis.tick_params(labelsize=8)
    pages.figure.text(.07, .065, 'Only supported intervals contribute to differences; gaps remain unknown.',
                      fontsize=8, color=MUTED, parse_math=False)


def _workspace_plots(pages, store, workspace, colors, show_cursors=False):
    from app_fast_plot import EnvelopeCache
    from app_time_units import minute_factor
    for pane_number, view in enumerate(visible_views(workspace), 1):
        if not view['ids']:
            pages.page('2. Graphs', f'Pane {pane_number}')
            pages.paragraph('No channels selected in this pane.')
            continue
        # Three independent axes and at most five traces keep print labels readable.
        groups=[];current=[];units=[]
        for id in view['ids']:
            channel=store.channels[id]
            key='Normalized offset' if view.get('normalized') else view.get('bindings', {}).get(id, channel.unit)
            if current and (len(current)>=5 or (key not in units and len(units)>=3)):
                groups.append(current);current=[];units=[]
            current.append(id)
            if key not in units:units.append(key)
        groups.append(current)
        for group_index, ids in enumerate(groups, 1):
            pages.page('2. Graphs', f'Pane {pane_number} · original workspace coordinates · part {group_index}/{len(groups)}')
            axis=pages.figure.add_axes([.12, .23, .67, .60])
            axes={};handles=[];labels=[]
            for id in ids:
                channel=store.channels[id]
                key='Normalized offset' if view.get('normalized') else view.get('bindings', {}).get(id, channel.unit)
                if key not in axes:
                    current_axis=axis if not axes else axis.twinx()
                    if len(axes)>1:current_axis.spines['right'].set_position(('outward', 45*(len(axes)-1)))
                    current_axis.set_ylabel(key or 'Value', color=colors[id], fontsize=9,parse_math=False)
                    current_axis.tick_params(axis='y', colors=colors[id], labelsize=8)
                    bounds=view.get('axes', {}).get(key)
                    if bounds:current_axis.set_ylim(*bounds)
                    elif not axes and view.get('range'):current_axis.set_ylim(*view['range'][1])
                    if key=='Actuators':current_axis.set_yticks([0,1], ['OFF','ON'])
                    axes[key]=current_axis
                x,y=channel.values()
                bounds=view.get('range', [(float(x[0]),float(x[-1])) if len(x) else (0,1)])[0]
                px,py=EnvelopeCache(x,y).visible(*bounds, 6000)
                if view.get('normalized'):
                    finite=y[np.isfinite(y)]
                    low,high=(float(finite.min()),float(finite.max())) if len(finite) else (0.,1.)
                    py=np.where(np.isfinite(py),(py-low)/(high-low) if high>low else .5,np.nan)+view['ids'].index(id)
                factor=minute_factor(channel.coordinate_unit) if channel.coordinate_kind=='time' else 1
                style=channel.attributes.get('plot_style', {})
                line,=axes[key].plot(px*factor, py, color=colors[id], label=channel_label(store,id),
                    drawstyle='steps-post' if channel.attributes.get('actuator_binary') else 'default',
                    linewidth=float(style.get('width', .8)),
                    linestyle={'solid':'-','dash':'--','dot':':','dashdot':'-.'}.get(style.get('style'),'-'))
                handles.append(line);labels.append(textwrap.shorten(channel_label(store,id),width=85,placeholder='…'))
                axis.set_xlim(bounds[0]*factor,bounds[1]*factor)
                axis.set_xlabel('Time (min)' if channel.coordinate_kind=='time' else
                               f'{channel.coordinate_kind} ({channel.coordinate_unit})',parse_math=False)
                if show_cursors and view.get('cursor_mode','dual')!='off':
                    for index, lock in enumerate(view.get('cursor_locks', [])[:2]):
                        if lock and lock.get('channel')==id and len(x):
                            sample=int(np.clip(lock['index'],0,len(x)-1))
                            value=y[sample]
                            if view.get('normalized'):
                                value=((value-low)/(high-low) if high>low else .5)+view['ids'].index(id)
                            axes[key].scatter([x[sample]*factor],[value], c=[CURSOR_COLORS[index]], s=35, zorder=6)
            primary=store.channels[ids[0]]
            factor=minute_factor(primary.coordinate_unit) if primary.coordinate_kind=='time' else 1
            if show_cursors and view.get('cursor_mode','dual')!='off':
                for index, position in enumerate(view.get('cursors', [])[:1 if view.get('cursor_mode')=='point' else 2]):
                    color=CURSOR_COLORS[index]
                    axis.axvline(position*factor,color=color,linestyle='--',linewidth=1.3)
                    axis.text(position*factor,.98,('A','B')[index]+f' {position*factor:.6g}',
                        color=color,backgroundcolor='#182033',transform=axis.get_xaxis_transform(),va='top',fontsize=8,parse_math=False)
            for annotation in view.get('annotations',[]):
                if annotation.get('kind')=='ROI':axis.axvspan(annotation['start']*factor,annotation['end']*factor,alpha=.15)
                elif annotation.get('kind')=='Threshold':axes.get(annotation.get('axis'),axis).axhline(annotation['start'],color='purple',linestyle=':')
                else:
                    axis.axvline(annotation['start']*factor,color='purple',linestyle=':')
                    if annotation.get('text'):axis.text(annotation['start']*factor,.88,annotation['text'],transform=axis.get_xaxis_transform(),fontsize=8,color='purple',parse_math=False)
            if view.get('show_grid',True):axis.grid(True,alpha=.18)
            if view.get('show_legends',True):
                legend=pages.figure.legend(handles, labels, loc='lower left', bbox_to_anchor=(.08,.075),
                    fontsize=7, framealpha=.45)
                for label in legend.get_texts():label.set_parse_math(False)
            pages.figure.text(.07,.16,'Plots retain the selected view bounds; cursor statistics and comparison scope are labelled separately.',
                              fontsize=7,color=MUTED,parse_math=False)


def render_comparison_pdf(path, store, workspace, evidence, summary, scorecard=None,
                          companion_names=None):
    from matplotlib.backends.backend_pdf import PdfPages
    from app_actuators import state_text
    pair=evidence.get('signal_comparison')
    report=evidence['report']
    metadata=report['Metadata']
    colors=_colors(store)
    with PdfPages(path, metadata={'Title':str(report['Report title']), 'Subject':'Four-section signal comparison evidence',
                                 'Keywords':'Report ID '+report['Report ID']}) as pdf:
        pages=_Pages(pdf,report)
        pages.page('1. Comparison overview')
        pages.paragraph(report['Report title'],label='Report title')
        pages.paragraph('Report ID: '+report['Report ID'])
        pages.paragraph('Statistics scope: '+str(report['Statistics scope']))
        if pair:
            pages.paragraph('Reference: '+text(pair.get('reference',{}).get('dataset'))+' · '+text(pair.get('reference',{}).get('name')))
            pages.paragraph('Candidate: '+text(pair.get('candidate',{}).get('dataset'))+' · '+text(pair.get('candidate',{}).get('name')))
            pages.paragraph(text(pair.get('alignment',{}).get('description')), label='Timeline alignment')
            for line in pair.get('summary') or [pair.get('reason') or 'No supported comparison result']:
                pages.paragraph(line)
            pages.table(_metric_rows(pair),
                        ['Measure','Value'], [.65,.35])
            pages.paragraph('Difference duration excludes floating-point roundoff; this is not an acceptance tolerance.',size=8)
            pages.paragraph('Differences are factual measurements, not a Pass/Fail decision. No acceptance limits are inferred.',size=8)
        else:
            pages.paragraph('No cross-file pair has been selected. This report provides each selected channel’s values and statistics; similar names do not establish a valid engineering pairing.')
        for key in ('Prepared by','Reviewed by','Test conditions'):
            pages.paragraph(metadata.get(key) or 'Not provided',label=key)
        pages.page('1. Comparison overview','Source files and scope')
        pages.table([{'Dataset':row['dataset'],'Source':row['source'],'Channels':row['selected_channels']} for row in evidence['sources']],
                    ['Dataset','Source','Channels'],[.28,.60,.12])
        pages.paragraph('Source-file hashes identify bytes available when this report was generated. Loaded-channel fingerprints identify the actual arrays and processing state used in this report.',size=8)
        for row in evidence['sources']:
            pages.paragraph(row['dataset'],label='Source fingerprint')
            pages.paragraph(row['fingerprint_method'],size=8)
            pages.paragraph(row.get('source_sha256') or 'Source file SHA-256 not available',size=8)
        if pair and pair.get('status')=='compared':_paired_plot(pages,pair,store,colors)
        from app_report_appearance import show_cursor_guides
        _workspace_plots(pages,store,workspace,colors,show_cursor_guides(metadata))
        pages.page('3. Measurements and differences','All selected channel intervals')
        pages.paragraph('A/B and B−A follow each pane’s actual cursor positions and original sample locks. Min/Max/Average scope follows the selected statistics scope. Coverage and RMS retain the active-interval calculation policy; timing quality is listed separately.',size=8)
        for row in summary.to_dict('records'):
            pages.ensure(.16)
            pages.paragraph(row['channel']+' ['+str(row['unit'])+']',label='Pane '+str(row['pane']))
            values={key:row.get(key) for key in ('A','B','ΔY','Min','Max','Mean','RMS','Coverage')}
            if row.get('value_interpretation')=='actuator':
                for key in ('A','B','Min','Max'):values[key]=state_text(values[key])
                values['ΔY']=state_text(row.get('A'))+' → '+state_text(row.get('B'))
                values['Mean']=f"{100*row['Mean']:.6g}% ON" if pd.notna(row.get('Mean')) else None
                values['RMS']='Not applicable'
            pages.table([values],list(values),[.125]*8)
            pages.paragraph(f"A coordinate {text(row.get('cursor_A_X'))}; B coordinate {text(row.get('cursor_B_X'))}; ΔX {text(row.get('ΔX'))} {row.get('coordinate_unit','')}",size=8)
            locks=[f"{name}: {row[name]}" for name in row if name.endswith(('locked_channel','sample_index')) and pd.notna(row[name])]
            if locks:pages.paragraph('; '.join(locks),size=7.5)
            score=row.get('match_percent')
            color=match_color(score)
            if color:
                pages.ensure(.025)
                pages.figure.text(.07,pages.y,f'Match {float(score):.1f}%',fontsize=8,color=INK,
                    bbox={'facecolor':color,'edgecolor':'none','alpha':.5},parse_math=False)
                pages.y-=.03
        if scorecard is not None and len(scorecard):
            pages.page('3. Measurements and differences','Existing configured comparison results')
            pages.paragraph('The following scorecard was supplied by the existing comparison workflow; no new limits or pairings were inferred.',size=8)
            for record in scorecard.to_dict('records'):
                pages.paragraph('; '.join(str(key)+': '+text(value) for key,value in record.items()),size=8)
        if pair:
            pages.page('3. Measurements and differences','Difference intervals and activity')
            statistics=[('Minimum','min'),('Maximum','max'),('Supported average','mean'),
                ('Supported RMS','rms'),('Supported standard deviation','std'),
                ('Recorded samples','samples'),('Observed span (min)','observed_span_min'),
                ('Supported duration (min)','supported_duration_min'),('Supported coverage (%)','coverage_fraction')]
            pages.paragraph('Each source uses its own supported intervals; difference metrics use intervals supported by both sources.',size=8)
            pages.table([{'Measure':label,**{role.title():
                (100*pair.get(role,{}).get('statistics',{}).get(key) if key=='coverage_fraction' and
                 pair.get(role,{}).get('statistics',{}).get(key) is not None else
                 pair.get(role,{}).get('statistics',{}).get(key)) for role in ('reference','candidate')}}
                for label,key in statistics],['Measure','Reference','Candidate'],[.50,.25,.25])
            for role in ('reference','candidate'):
                source=pair.get(role,{})
                pages.paragraph(role.title()+': '+text(source.get('dataset'))+' · '+text(source.get('name')),label='Source activity')
                names={'starts':'Zero → nonzero transitions','stops':'Nonzero → zero transitions',
                    'sign_changes':'Numeric sign changes (not a physical interpretation)',
                    'reversals':'Numeric sign changes (not a physical interpretation)',
                    'positive_duration_min':'Supported positive-reading duration (min)',
                    'negative_duration_min':'Supported negative-reading duration (min)'}
                activity=source.get('activity',{})
                pages.table([{'Measure':names.get(key,key.replace('_',' ')),'Value':text(value)} for key,value in activity.items()
                             if key!='reversals' or 'sign_changes' not in activity],
                            ['Measure','Value'],[.65,.35])
            records=pair.get('differences') or []
            supported=[row for row in records if row.get('supported')]
            samples=supported[:6]+[row for row in records if not row.get('supported')][:6]
            pages.paragraph(f'{len(records):,} complete difference intervals and {len(pair.get("events") or []):,} activity events are retained in workbook/JSON. This PDF shows {len(samples)} representative supported/unknown intervals; it does not omit evidence from the companion files.',size=8)
            if samples:
                pages.table([{'Start min':row.get('start_min'),'End min':row.get('end_min'),
                              'Δ start':row.get('delta_start'),'Δ end':row.get('delta_end'),
                              'Support':row.get('reason') or ('Supported' if row.get('supported') else 'Unknown')} for row in samples],
                            ['Start min','End min','Δ start','Δ end','Support'],[.17,.17,.16,.16,.34])
        pages.page('4. Quality, traceability and conclusions')
        pages.paragraph('Data gaps, missing values and unreviewed acquisition assumptions limit conclusions. Data quality is separate from appliance faults and from configured acceptance criteria.',size=8)
        for row in evidence['channel_quality']:
            pages.paragraph(row['channel'],label='Pane '+str(row['pane'])+' · acquisition quality')
            selected={key:row.get(key) for key in ('recorded_samples','missing_readings','duplicate_timestamps',
                'gap_count','interval_policy','quality_flags') if key in row}
            for key,label in [('supported_duration_s','supported duration (min)'),
                              ('observed_span_s','observed span (min)'),('gap_threshold_s','gap threshold (min)')]:
                if key in row:selected[label]=row[key]/60 if row[key] is not None else None
            if row.get('coverage') is not None:selected['supported coverage (%)']=100*row['coverage']
            pages.paragraph('; '.join(key.replace('_',' ')+': '+text(value) for key,value in selected.items()),size=8)
        for limitation in (pair or {}).get('limitations') or []:pages.paragraph(limitation,size=8)
        for key in ('Observations','Conclusions','Recommended actions'):
            pages.paragraph(metadata.get(key) or 'Not provided',label='User-provided '+key.lower())
        pages.paragraph('Manual notes do not change computed values or acceptance rules.',size=8)
        pages.page('4. Quality, traceability and conclusions','Loaded measurements and processing')
        for row in evidence['processing_provenance']:
            pages.paragraph(row['channel'],label='Loaded channel snapshot')
            pages.paragraph(row['loaded_snapshot_sha256'],size=8)
            pages.paragraph('Source column: '+str(row['source_column'])+'; shift: '+str(row['shift'])+' '+str(row['coordinate_unit'])+
                '; masks: '+text(row['masks'])+'; parents: '+text(row['parents'])+'; recipe: '+text(row['recipe']),size=8)
        for name in companion_names or []:pages.paragraph(name,label='Complete evidence file')
        pages.save()
        return pages.page_number


def export_comparison_bundle(path, store, workspace, metadata=None, scorecard=None,
                             *, title='Signal comparison report', cancel=None,_prepared=None):
    """Write the readable PDF, full workbook and strict JSON as one transaction."""
    from app_reporting import interval_tables
    from app_engineering_reports import staged_report_bundle
    path=Path(path).with_suffix('.pdf')
    if not path.parent.is_dir():raise ValueError('Choose an existing writable output folder')
    workbook=path.with_suffix('.xlsx');json_path=path.with_suffix('.evidence.json')
    checkpoint('Collecting signal comparison evidence')
    data,summary=_prepared if _prepared is not None else interval_tables(store,workspace)
    evidence,tables=build_evidence(store,workspace,data,summary,metadata,scorecard,title)
    destinations=[path,workbook,json_path]
    with staged_report_bundle(destinations,cancel) as staged:
        write_workbook(staged[workbook],data,summary,evidence,tables,scorecard)
        pages=render_comparison_pdf(staged[path],store,workspace,evidence,summary,scorecard,
                                   [workbook.name,json_path.name])
        evidence['pdf_pages']=pages
        evidence['companions']=[item.name for item in destinations]
        checkpoint('Writing complete comparison JSON evidence')
        staged[json_path].write_text(json.dumps(clean(evidence),indent=2,ensure_ascii=False,allow_nan=False)+'\n',encoding='utf-8')
        checkpoint('Committing complete comparison report')
    return {'pdf':str(path),'workbook':str(workbook),'evidence':str(json_path),
            'report_id':evidence['report']['Report ID']}
