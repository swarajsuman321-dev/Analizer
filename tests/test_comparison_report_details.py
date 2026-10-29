"""User-facing report completeness, agreement and failure-safety checks."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd

from app_channel_store import Channel, ChannelStore, Dataset
from app_comparison_report import build_evidence, export_comparison_bundle, write_workbook, worksheet_parts
from app_reporting import export_report, interval_tables


class ComparisonReportTests(unittest.TestCase):
    def setUp(self):
        self.temporary=tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root=Path(self.temporary.name)
        self.store=ChannelStore()
        self.ids=[]
        for index,name in enumerate(('Reference $\\invalid$', 'Candidate')):
            source=self.root/(name.replace('\\','')+'.csv')
            source.write_text('Time (min),Speed\n0,0\n1,10\n2,20\n')
            dataset=Dataset(name,pd.DataFrame({'Time':[0,1,2],'Speed':[0,10,20]}),str(source))
            self.store.datasets[dataset.id]=dataset
            channel=Channel('Drum speed',[0,1,1,2,20],[0,10+index,11+index,np.nan,4],
                unit='RPM',coordinate_unit='min',dataset_id=dataset.id,source=str(source),
                attributes={'plot_style':{'color':('#075985','#b45309')[index]},'match_percent':93})
            self.store.add(channel);self.ids.append(channel.id)
        self.workspace={'layout':'Horizontal split','panes':[self.view(id) for id in self.ids],
            'report_kind':'comparison','workspace_name':'Comparison','statistics_scope':'Cursor interval'}
        self.notes={'Report title':'Reviewed comparison $\\not_math$',
            'Prepared by':'Engineer A','Reviewed by':'Engineer B',
            'Test conditions':'Same load, independent raw recordings',
            'Observations':'Observed difference, not an acceptance decision.',
            'Conclusions':'Review the supported overlap.',
            'Recommended actions':'Confirm sampling and rerun the selected interval.'}

    def view(self,id):
        return {'ids':[id],'cursors':[1,2],'cursor_locks':[{'channel':id,'index':1},None],
            'snap':True,'bindings':{id:'RPM'},'axes':{'RPM':[-1,22]},
            'range':[[0,20],[-1,22]],'cursor_mode':'dual','show_legends':True}

    def pair(self):
        sources={}
        for role,id in zip(('reference','candidate'),self.ids):
            channel=self.store.channels[id]
            sources[role]={'channel_id':id,'name':channel.name,'dataset':self.store.datasets[channel.dataset_id].name,
                'unit':'RPM','statistics':{'mean':10,'rms':12},'quality':{'missing_readings':1},
                'activity':{'available':True,'interpretation':'Exact nonzero speed','starts':1,'stops':0},
                'plot':{'time_min':[0,1,2],'value':[0,10,20]}}
        return {'schema':1,'status':'compared',**sources,
            'alignment':{'mode':'elapsed','description':'Elapsed from each recording start'},
            'scope':{'start_min':0,'end_min':20},
            'metrics':{'unit':'RPM','mean_bias':1.1234567890123,'rmse':1.,'compared_duration_min':1.,'coverage_fraction':.355},
            'differences':[{'start_min':i,'end_min':i+1,'duration_min':1.,'supported':i%2==0,
                'reference_start':0,'reference_end':10,'candidate_start':1,'candidate_end':11,
                'delta_start':1. if i%2==0 else None,'delta_end':1. if i%2==0 else None,
                'reason':'Supported' if i%2==0 else 'Missing source support'} for i in range(30)],
            'events':[{'role':'candidate','event':'start','time_min':i,'from_value':0,'to_value':1} for i in range(18)],
            'summary':['Candidate average exceeds reference by 1 RPM on supported intervals.'],
            'limitations':['Unknown gaps are excluded.']}

    def test_bundle_contains_all_evidence_and_exact_locked_sample_without_source_mutation(self):
        metadata={**self.notes,'Signal comparison':self.pair()}
        attrs=deepcopy({id:c.attributes for id,c in self.store.channels.items()})
        arrays={id:(c.x.copy(),c.y.copy()) for id,c in self.store.channels.items()}
        result=export_comparison_bundle(self.root/'Comparison.pdf',self.store,self.workspace,metadata)
        evidence=json.loads(Path(result['evidence']).read_text())
        sheets=pd.read_excel(result['workbook'],sheet_name=None)
        self.assertEqual(evidence['schema'],1)
        self.assertEqual(evidence['report']['Report ID'],result['report_id'])
        self.assertEqual(evidence['report']['Metadata']['Conclusions'],self.notes['Conclusions'])
        self.assertEqual(len(evidence['signal_comparison']['differences']),30)
        self.assertEqual(len(sheets['Difference intervals']),30)
        self.assertEqual(len(sheets['Activity events']),18)
        self.assertAlmostEqual(float(sheets['Pair comparison'].set_index('Field').loc['metrics.mean_bias','Value']),1.1234567890123,places=12)
        self.assertEqual(sheets['Cursor summary'].iloc[0]['A'],10)
        self.assertEqual(evidence['cursor_summary'][0]['A'],10)
        self.assertEqual(evidence['cursor_summary'][0]['cursor_A_sample_index'],1)
        self.assertEqual(sheets['Cursor summary'].iloc[0]['match_percent'],93)
        self.assertEqual(sheets['Channel quality'].iloc[0]['duplicate_timestamps'],1)
        for row in evidence['sources']:
            self.assertEqual(row['source_sha256'],hashlib.sha256(Path(row['source']).read_bytes()).hexdigest())
            self.assertIn('not an import-time identity claim',row['fingerprint_method'])
        self.assertEqual(evidence['report']['Report ID'],sheets['Report details'].set_index('Field').loc['Report ID','Value'])
        self.assertEqual(Path(result['pdf']).read_bytes()[:5],b'%PDF-')
        self.assertEqual([row['trace_color'] for row in evidence['processing_provenance']],
                         ['#075985','#b45309'])
        self.assertEqual(len(re.findall(rb'/Type\s*/Page\b',Path(result['pdf']).read_bytes())),evidence['pdf_pages'])
        self.assertEqual(attrs,{id:c.attributes for id,c in self.store.channels.items()})
        for id,c in self.store.channels.items():
            np.testing.assert_equal(c.x,arrays[id][0]);np.testing.assert_equal(c.y,arrays[id][1])

    def test_elapsed_overlay_uses_the_selected_channel_colors(self):
        from matplotlib.backends.backend_pdf import PdfPages
        original=PdfPages.savefig
        paired=[]
        def inspect(pdf,figure,*args,**kwargs):
            if any('Elapsed comparison of the explicitly selected signals' in item.get_text()
                   for item in figure.texts):
                paired.extend(line.get_color() for line in figure.axes[0].lines)
            return original(pdf,figure,*args,**kwargs)
        self.store.channels[self.ids[0]].attributes['plot_style']['color']='#993399'
        self.store.channels[self.ids[1]].attributes['plot_style']['color']='#339933'
        with patch.object(PdfPages,'savefig',inspect):
            export_comparison_bundle(self.root/'Colors.pdf',self.store,self.workspace,
                                     {**self.notes,'Signal comparison':self.pair()})
        self.assertEqual(paired,['#993399','#339933'])

    def test_numerical_resolution_evidence_does_not_expand_the_pdf_summary(self):
        from matplotlib.backends.backend_pdf import PdfPages
        from matplotlib.text import Text
        pair=self.pair()
        policy='Full numerical resolution definition: '+('conversion and interpolation roundoff evidence; '*8)
        pair['metrics'].update(numerical_resolution_max=3.552713678800501e-14,
                               numerical_resolution_policy=policy,
                               different_duration_min=0.,
                               difference_direction='Candidate minus Reference',
                               weighting='Elapsed-time weighted joint support')
        first_page=[]
        original=PdfPages.savefig
        def inspect(pdf,figure,*args,**kwargs):
            if not first_page:
                first_page.extend(item.get_text() for item in figure.findobj(Text))
            return original(pdf,figure,*args,**kwargs)
        with patch.object(PdfPages,'savefig',inspect):
            result=export_comparison_bundle(self.root/'Resolution.pdf',self.store,self.workspace,
                                            {**self.notes,'Signal comparison':pair})
        joined='\n'.join(first_page)
        self.assertIn('Supported time with distinguishable differences',joined)
        self.assertIn('Difference duration excludes floating-point roundoff',joined)
        self.assertNotIn('numerical_resolution',joined)
        self.assertNotIn(policy,joined)
        self.assertNotIn('Configured comparison tolerance',joined)
        evidence=json.loads(Path(result['evidence']).read_text())
        self.assertEqual(evidence['signal_comparison']['metrics']['numerical_resolution_policy'],policy)
        metrics=pd.read_excel(result['workbook'],sheet_name='Pair comparison').set_index('Field')['Value']
        self.assertEqual(metrics['metrics.numerical_resolution_policy'],policy)
        self.assertEqual(float(metrics['metrics.numerical_resolution_max']),3.552713678800501e-14)

    def test_real_pdf_pages_remain_standard_size_and_include_every_channel_and_notes(self):
        extra=[]
        for index in range(13):
            channel=self.store.add(Channel('Extra parameter '+str(index),[0,1,2],[index,index+1,index+2],
                unit='RPM',coordinate_unit='min',dataset_id=self.store.channels[self.ids[0]].dataset_id))
            extra.append(channel.id)
        self.workspace['panes'][0]['ids'].extend(extra)
        from matplotlib.backends.backend_pdf import PdfPages
        from matplotlib.text import Text
        original=PdfPages.savefig
        pages=[];content=[]
        def inspect(pdf,figure,*args,**kwargs):
            pages.append(tuple(figure.get_size_inches()))
            content.extend(item.get_text() for item in figure.findobj(Text))
            return original(pdf,figure,*args,**kwargs)
        with patch.object(PdfPages,'savefig',inspect):
            result=export_comparison_bundle(self.root/'Complete.pdf',self.store,self.workspace,self.notes)
        self.assertGreater(len(pages),6)
        self.assertTrue(all(np.allclose(size,[8.27,11.69]) for size in pages))
        joined='\n'.join(content)
        for index in range(13):self.assertIn('Extra parameter '+str(index),joined)
        for section in ('1. Comparison overview','2. Graphs','3. Measurements and differences','4. Quality, traceability and conclusions'):
            self.assertIn(section,joined)
        for key in ('Observations','Conclusions','Recommended actions'):
            self.assertIn(self.notes[key],joined)
        self.assertIn('No cross-file pair has been selected.',joined)
        self.assertEqual(json.loads(Path(result['evidence']).read_text())['raw_interval_sample_count'],
                         len(pd.read_excel(result['workbook'],sheet_name='Active interval')))

    def test_comparison_export_routes_to_bundle_and_xlsx_retains_existing_sheets(self):
        pdf=self.root/'Quick.pdf'
        data,summary=export_report(pdf,self.store,self.workspace,metadata=self.notes)
        self.assertTrue(pdf.with_suffix('.xlsx').exists())
        self.assertTrue(pdf.with_suffix('.evidence.json').exists())
        xlsx=self.root/'Only workbook.xlsx'
        export_report(xlsx,self.store,self.workspace,metadata=self.notes)
        sheets=pd.read_excel(xlsx,sheet_name=None)
        self.assertTrue({'Active interval','Cursor summary','Sources','Report details','Channel quality','View settings','Processing provenance'}<=set(sheets))
        self.assertEqual(len(sheets['Active interval']),len(data))
        self.assertEqual(len(sheets['Cursor summary']),len(summary))
        self.assertTrue(xlsx.with_suffix('.evidence.json').exists())

    def test_analysis_and_csv_keep_original_export_contract(self):
        analysis=deepcopy(self.workspace);analysis.pop('report_kind');analysis['workspace_name']='Analysis'
        with patch('app_workspace_io.render_report') as render:
            export_report(self.root/'Analysis.pdf',self.store,analysis,metadata=self.notes)
        render.assert_called_once()
        self.assertFalse((self.root/'Analysis.xlsx').exists())
        csv=self.root/'Samples.csv'
        expected,_=interval_tables(self.store,self.workspace)
        export_report(csv,self.store,self.workspace,metadata=self.notes)
        actual=pd.read_csv(csv)
        self.assertEqual(list(actual.columns),list(expected.columns))
        np.testing.assert_equal(actual.coordinate.to_numpy(),expected.coordinate.to_numpy())
        self.assertEqual(len(actual),len(expected))

    def test_unavailable_sources_and_non_temporal_quality_do_not_invent_time_coverage(self):
        channel=self.store.add(Channel('Spectrum',[0,1,2],[1,np.nan,3],unit='W',coordinate_unit='Hz',coordinate_kind='frequency'))
        workspace={'layout':'Single','panes':[{'ids':[channel.id],'cursors':[0,2]}],'report_kind':'comparison'}
        data,summary=interval_tables(self.store,workspace)
        evidence,tables=build_evidence(self.store,workspace,data,summary,self.notes)
        row=evidence['channel_quality'][0]
        self.assertNotIn('coverage',row)
        self.assertIn('not applicable',row['interval_policy'])
        self.assertIsNone(evidence['sources'][0]['source_sha256'])
        self.assertTrue(evidence['processing_provenance'][0]['loaded_snapshot_sha256'])
        self.assertIsNone(evidence['cursor_summary'][0].get('Mean'))
        json.dumps(evidence,allow_nan=False)

    def test_long_notes_are_retained_in_extended_text_not_silently_clipped_by_excel(self):
        note='Start '+('large reviewed observation '*1600)+' End'
        data,summary=interval_tables(self.store,self.workspace)
        evidence,tables=build_evidence(self.store,self.workspace,data,summary,{'Observations':note})
        path=self.root/'Long.xlsx'
        write_workbook(path,data,summary,evidence,tables)
        sheets=pd.read_excel(path,sheet_name=None)
        rows=sheets['Extended text'].sort_values('Part')
        self.assertEqual(''.join(rows.Text),note)
        self.assertEqual(evidence['report']['Metadata']['Observations'],note)

    def test_workbook_notes_and_source_names_remain_literal_text(self):
        from openpyxl import load_workbook
        self.store.datasets[self.store.channels[self.ids[0]].dataset_id].name='=Appliance A'
        notes={'Observations':'=SUM(1,2)','Conclusions':'#REF!'}
        data,summary=interval_tables(self.store,self.workspace)
        evidence,tables=build_evidence(self.store,self.workspace,data,summary,notes)
        path=self.root/'Literal.xlsx'
        write_workbook(path,data,summary,evidence,tables)
        book=load_workbook(path,data_only=False)
        try:
            values={row[0].value:row[1] for row in book['Report details'].iter_rows(min_row=2)}
            self.assertEqual(values['Metadata.Observations'].value,'=SUM(1,2)')
            self.assertEqual(values['Metadata.Observations'].data_type,'s')
            self.assertEqual(values['Metadata.Conclusions'].value,'#REF!')
            self.assertEqual(values['Metadata.Conclusions'].data_type,'s')
        finally:book.close()
        sheets=pd.read_excel(path,sheet_name=None)
        self.assertEqual(sheets['Report details'].set_index('Field').loc['Metadata.Observations','Value'],'=SUM(1,2)')
        self.assertEqual(sheets['Sources'].iloc[0]['dataset'],'=Appliance A')

    def test_analysis_workbook_retains_literal_names_and_numeric_readings(self):
        from openpyxl import load_workbook
        self.store.channels[self.ids[0]].name='=SUM(1,2)'
        analysis=deepcopy(self.workspace);analysis['report_kind']='analysis';analysis['workspace_name']='Analysis'
        path=self.root/'Analysis.xlsx'
        export_report(path,self.store,analysis)
        book=load_workbook(path,data_only=False)
        try:
            sheet=book['Active interval']
            columns={cell.value:cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2,columns['original_channel']).value,'=SUM(1,2)')
            self.assertEqual(sheet.cell(2,columns['original_channel']).data_type,'s')
            self.assertEqual(sheet.cell(2,columns['value']).value,10)
            self.assertEqual(sheet.cell(2,columns['value']).data_type,'n')
        finally:book.close()
        self.assertEqual(pd.read_excel(path,sheet_name='Active interval').iloc[0]['original_channel'],'=SUM(1,2)')
        self.assertFalse(path.with_suffix('.evidence.json').exists())

    def test_large_table_partition_plan_and_actual_small_partition_keep_every_row(self):
        self.assertEqual(worksheet_parts('Difference intervals',1048575),[('Difference intervals',0,1048575)])
        self.assertEqual(worksheet_parts('Difference intervals',1048576),
            [('Difference intervals',0,1048575),('Difference intervals 2',1048575,1048576)])
        self.assertTrue(all(len(name)<=31 for name,_,_ in worksheet_parts('x'*31,300,100)))
        data,summary=interval_tables(self.store,self.workspace)
        data=pd.concat([data]*4,ignore_index=True)
        evidence,tables=build_evidence(self.store,self.workspace,data,summary,self.notes)
        path=self.root/'Partitions.xlsx'
        with patch('app_comparison_report.EXCEL_DATA_ROWS',5):
            write_workbook(path,data,summary,evidence,tables)
        sheets=pd.read_excel(path,sheet_name=None)
        names=[name for name in sheets if name=='Active interval' or name.startswith('Active interval ')]
        combined=pd.concat([sheets[name] for name in names],ignore_index=True)
        self.assertEqual(len(combined),len(data))
        np.testing.assert_equal(combined.coordinate.to_numpy(),data.coordinate.to_numpy())
        self.assertTrue(all(len(sheets[name])<=5 for name in names))

    def test_cancelled_bundle_preserves_all_previous_artifacts(self):
        pdf=self.root/'Previous.pdf';paths=[pdf,pdf.with_suffix('.xlsx'),pdf.with_suffix('.evidence.json')]
        for path in paths:path.write_bytes(('old '+path.suffix).encode())
        before={path:path.read_bytes() for path in paths}
        with patch('app_comparison_report.render_comparison_pdf',side_effect=InterruptedError('cancelled')):
            with self.assertRaises(InterruptedError):
                export_comparison_bundle(pdf,self.store,self.workspace,self.notes)
        self.assertEqual(before,{path:path.read_bytes() for path in paths})
        self.assertFalse(list(self.root.glob('.signal-report-*')))

    def test_failed_commit_restores_entire_previous_set(self):
        pdf=self.root/'Previous.pdf';paths=[pdf,pdf.with_suffix('.xlsx'),pdf.with_suffix('.evidence.json')]
        for path in paths:path.write_bytes(('old '+path.suffix).encode())
        before={path:path.read_bytes() for path in paths}
        import os
        original=os.replace;count=0
        def fail_once(source,destination):
            nonlocal count
            count+=1
            if count==5:raise OSError('injected publication failure')
            return original(source,destination)
        with patch('app_engineering_reports.os.replace',side_effect=fail_once):
            with self.assertRaisesRegex(OSError,'injected'):
                export_comparison_bundle(pdf,self.store,self.workspace,self.notes)
        self.assertEqual(before,{path:path.read_bytes() for path in paths})
        self.assertFalse(list(self.root.glob('.signal-report-*')))


if __name__=='__main__':unittest.main()
