"""Complete spreadsheet evidence with bounded writer memory and cancellation."""
from datetime import date, datetime, timedelta
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES

from app_comparison_report import write_workbook


class StreamingComparisonWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.path = Path(self.temporary.name)/'Evidence.xlsx'

    def test_complete_raw_and_difference_tables_keep_native_precision_and_literal_text(self):
        count = 10003
        frame = pd.DataFrame({'sample':np.arange(count),
            'value':np.arange(count,dtype=float)*1.234567890123e-8,
            'name':['=SUM(1,2)']*(count-1)+['#REF!'],
            'missing':[np.nan]*count})
        differences = pd.DataFrame({'start_min':np.arange(count,dtype=float)/60,
            'delta':np.arange(count,dtype=float)*-1.234567890123e-12,
            'supported':np.arange(count)%2==0})
        evidence = {}
        write_workbook(self.path,frame,pd.DataFrame({'match_percent':[93.,75.,20.]}),
                       evidence,{'Difference intervals':differences})
        book = load_workbook(self.path,read_only=True,data_only=False)
        try:
            # Published dimensions let independent consumers count every row
            # without loading millions of Cell objects or inferring truncation.
            self.assertEqual(book['Active interval'].max_row,count+1)
            self.assertEqual(book['Difference intervals'].max_row,count+1)
            for expected,row in zip(frame.itertuples(index=False,name=None),
                                    book['Active interval'].iter_rows(min_row=2)):
                self.assertEqual(row[0].value,expected[0])
                self.assertAlmostEqual(row[1].value,expected[1],places=16)
                self.assertEqual(row[2].value,expected[2])
                self.assertEqual(row[2].data_type,'s')
                self.assertEqual(row[3].value,'Not Available')
            for expected,row in zip(differences.itertuples(index=False,name=None),
                                    book['Difference intervals'].iter_rows(min_row=2)):
                self.assertAlmostEqual(row[1].value,expected[1],places=20)
                self.assertIs(row[2].value,expected[2])
        finally:
            book.close()
        self.assertTrue(any(row['Data rows']==count for row in evidence['workbook_parts']))

    def test_headers_filters_match_accents_and_dates_remain_available(self):
        frame = pd.DataFrame({'reading':[1.,2.,3.],
            'match_percent':[93.,75.,20.],
            'timestamp':[datetime(2026,10,29,12,34,56)]*3,
            'day':[date(2026,10,29)]*3,
            'elapsed':[timedelta(seconds=90)]*3})
        write_workbook(self.path,frame,pd.DataFrame({'empty':[]}),{}, {})
        book = load_workbook(self.path,data_only=False)
        try:
            sheet=book['Active interval']
            self.assertEqual(sheet.freeze_panes,'A2')
            self.assertEqual(sheet.auto_filter.ref,'A1:E4')
            self.assertTrue(sheet['A1'].font.bold)
            self.assertEqual(sheet['A1'].fill.fgColor.rgb,'00243A53')
            self.assertEqual([sheet.cell(row,2).fill.fgColor.rgb for row in range(2,5)],
                             ['0022a06b','00e5a000','00dc4545'])
            self.assertEqual(sheet['C2'].value,datetime(2026,10,29,12,34,56))
            self.assertEqual(sheet['D2'].value.date(),date(2026,10,29))
            self.assertAlmostEqual(sheet['E2'].value,90/86400)
        finally:
            book.close()

    def test_cancellation_removes_streaming_temporary_xml(self):
        before=set(ALL_TEMP_FILES)
        def checkpoint(stage=None,current=None,total=None):
            if current is not None and current>=4096:
                raise InterruptedError('cancel requested')
        with patch('app_comparison_report.checkpoint',checkpoint):
            with self.assertRaisesRegex(InterruptedError,'cancel requested'):
                write_workbook(self.path,pd.DataFrame({'value':np.arange(5000)}),
                               pd.DataFrame(),{}, {})
        self.assertFalse(self.path.exists())
        self.assertEqual(set(ALL_TEMP_FILES),before)

    def test_empty_and_partitioned_tables_keep_all_records_and_complete_long_text(self):
        note='=review '+('complete observation '*1800)+' #REF!'
        frame=pd.DataFrame({'sample':list(range(9)),'note':[note]+['ok']*8})
        with patch('app_comparison_report.EXCEL_DATA_ROWS',4):
            write_workbook(self.path,frame,pd.DataFrame({'empty':[]}),{}, {})
        book=load_workbook(self.path,read_only=True,data_only=False)
        try:
            names=['Active interval','Active interval 2','Active interval 3']
            samples=[row[0] for name in names for row in book[name].iter_rows(min_row=2,values_only=True)]
            self.assertEqual(samples,list(range(9)))
            parts=[row for name in book.sheetnames if name.startswith('Extended text')
                   for row in book[name].iter_rows(min_row=2,values_only=True)]
            self.assertEqual(''.join(row[-1] for row in sorted(parts,key=lambda row:row[-2])),note)
            self.assertEqual(book['Cursor summary'].max_row,1)
        finally:
            book.close()


if __name__=='__main__':unittest.main()
