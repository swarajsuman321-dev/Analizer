"""Reports retain all comparison/source evidence without unreadable PDF pages."""
from contextlib import contextmanager
import json
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.text import Text
from matplotlib.table import Table

from app_appliance_profiles import load_profile
from app_channel_store import Channel, ChannelStore, Dataset
from app_engineering_analytics import analyze_run
from app_engineering_compare import compare_runs
from app_engineering_reports import export_engineering, report_preview, report_tables, SHEETS


@contextmanager
def rendered_pages():
    """Observe text artists while retaining the actual PDF backend/file output."""
    pages = []
    save = PdfPages.savefig
    def record(pdf, figure, *args, **kwargs):
        texts=[artist.get_text() for artist in figure.findobj(match=Text)]
        # Table cells deliberately do not expose their Text as ordinary child
        # artists, so include the cells actually passed to the PDF renderer.
        for table in figure.findobj(match=Table):
            texts.extend(cell.get_text().get_text() for cell in table.get_celld().values())
        pages.append(texts)
        return save(pdf, figure, *args, **kwargs)
    with patch.object(PdfPages, 'savefig', record):
        yield pages


class EngineeringReportDetailsTests(unittest.TestCase):
    def fixture(self, count=4, comparison=True):
        store = ChannelStore()
        runs = []
        selected = []
        for number, role in enumerate(('reference', 'candidate') if comparison else ('analysis',)):
            dataset = Dataset(role.title()+' recording', pd.DataFrame(), '')
            store.datasets[dataset.id] = dataset
            for index in range(count):
                channel = Channel('Power' if index == 0 else f'Extra {role} parameter {index:02}',
                    np.arange(6.), np.arange(6.)*(index+1)+number*10.,
                    unit='W' if index == 0 else 'V', coordinate_unit='min', dataset_id=dataset.id,
                    source=f'{role}-source-{index:02}.csv')
                selected.append(store.add(channel).id)
            run = analyze_run(store, load_profile(), dataset_id=dataset.id)
            run['role'] = role
            runs.append(run)
        compared = compare_runs(runs[0], runs[1], runs[0]['profile']) if comparison else None
        return store, runs, compared, selected

    def test_all_comparison_rows_and_both_cursor_sources_reach_actual_pdf(self):
        store, runs, compared, selected = self.fixture()
        metadata = {'Prepared by': 'Review Engineer',
                    'Test conditions': 'Bench B, logger specification reviewed',
                    'Observations': 'Literal cost $5 and \\unparsed{value}; observation token OBS42.',
                    'Conclusions': 'User assessment only: CONCLUSION42.',
                    'Recommended actions': 'Review the configured test plan: ACTION42.'}
        # Distinct result names establish that even the final comparison row is
        # rendered rather than merely counted in a sidecar or PDF page plan.
        for index, row in enumerate(compared['kpis']):
            row['name'] = f'Comparison evidence metric {index:02}'
        before = [(channel.x.copy(), channel.y.copy()) for channel in store.channels.values()]
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'Detailed.pdf'
            with rendered_pages() as pages:
                tables = export_engineering(path, store, runs, compared, metadata,
                    view={'cursors': [1., 4.], 'snap': True}, selected_ids=selected)
            evidence = json.loads(path.with_suffix('.evidence.json').read_text())
            plan = evidence['page_plan']
            self.assertEqual(len(pages), plan['expected_pages'])
            self.assertEqual(len(re.findall(rb'/Type\s*/Page\b', path.read_bytes())), len(pages))
            comparison_pages = [page for page in pages if '3. Complete comparison results' in page]
            self.assertEqual(len(comparison_pages), plan['comparison_continuation_pages'])
            comparison_text = ' '.join('\n'.join(text for page in comparison_pages for text in page).split())
            for row in compared['kpis']:
                self.assertIn(row['name'], comparison_text)
            shared_reason='Existing KPI values compared for the stated scope; acceptance requires configured test-plan evidence'
            self.assertNotIn(shared_reason, comparison_text)
            self.assertIn('acceptance requires configured test-plan evidence', comparison_text)
            # The PDF avoids repeated boilerplate without dropping the original
            # explanation or any unavailable-result blocker from the evidence.
            comparison_records=tables['Comparison Results'].to_dict('records')
            self.assertTrue(any(row['result_state']=='Available' and row['comparison_reason']==shared_reason
                                for row in comparison_records))
            self.assertTrue(any(row['result_state']=='Not Available' and row['comparison_reason']
                                for row in comparison_records))
            cursors = evidence['metadata']['Cursor interval evidence']
            self.assertEqual(len(cursors), len(selected))
            self.assertEqual({row['run'] for row in cursors}, {'reference', 'candidate'})
            self.assertEqual({row['channel_id'] for row in cursors}, set(selected))
            cursor_pages = [page for page in pages if '3. Complete cursor interval evidence' in page]
            self.assertEqual(len(cursor_pages), plan['cursor_continuation_pages'])
            cursor_text = ' '.join('\n'.join(text for page in cursor_pages for text in page).split())
            for channel_id in selected:
                self.assertIn(store.channels[channel_id].name, cursor_text)
            all_text = ' '.join('\n'.join(text for page in pages for text in page).split())
            for token in ('OBS42', 'CONCLUSION42', 'ACTION42', '\\unparsed{value}', '$5'):
                self.assertIn(token, all_text)
            self.assertIn('user-provided notes', all_text)
            workbook = pd.read_excel(path.with_suffix('.evidence.xlsx'), sheet_name=None)
            self.assertEqual(list(workbook), SHEETS)
            self.assertEqual(len(workbook['Comparison Results']), len(compared['kpis']))
            self.assertEqual(len(tables['Comparison Results']), plan['comparison_records'])
            written=pd.to_numeric(workbook['Comparison Results']['candidate'],errors='coerce')
            original=pd.to_numeric(tables['Comparison Results']['candidate'],errors='coerce')
            np.testing.assert_array_equal(written.isna(),original.isna())
            np.testing.assert_allclose(written,original,rtol=1e-12,atol=1e-15,equal_nan=True)
            self.assertEqual(evidence['metadata']['Reviewed report notes']['Conclusions'], metadata['Conclusions'])
        for channel, (x, y) in zip(store.channels.values(), before):
            np.testing.assert_array_equal(channel.x, x)
            np.testing.assert_array_equal(channel.y, y)

    def test_selected_source_and_quality_continuations_keep_every_channel(self):
        store, runs, _, selected = self.fixture(count=15, comparison=False)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'All selected.pdf'
            with rendered_pages() as pages:
                export_engineering(path, store, runs, selected_ids=selected)
            evidence = json.loads(path.with_suffix('.evidence.json').read_text())
            plan = evidence['page_plan']
            self.assertEqual(plan['selected_plot_records'], len(selected))
            self.assertEqual(plan['quality_records'], len(selected))
            self.assertEqual(plan['source_records'], len(selected))
            self.assertGreater(plan['graph_continuation_pages'], 0)
            self.assertGreater(plan['quality_continuation_pages'], 0)
            self.assertGreater(plan['source_continuation_pages'], 0)
            self.assertEqual(len(pages), plan['expected_pages'])
            all_text = '\n'.join(text for page in pages for text in page)
            for channel in store.channels.values():
                self.assertIn(channel.name, all_text)
                self.assertIn(channel.source, all_text)
            self.assertEqual({row['channel_id'] for row in evidence['metadata']['Source timing quality']}, set(selected))

    def test_notes_continue_without_clipping_or_interpreting_acceptance(self):
        store, runs, _, _ = self.fixture(count=1, comparison=False)
        notes = {'Report title': 'Report '+('long title '*30),
                 'Observations': ('Reviewed observation with literal $price and \\unsupported{input}. '*100)+' FINALOBSERVATION42',
                 'Conclusions': 'Reviewer opinion; acceptance remains Not Evaluated.'}
        expected = report_preview(runs, store=store, metadata=notes)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'Notes.pdf'
            with rendered_pages() as pages:
                export_engineering(path, store, runs, metadata=notes)
            evidence = json.loads(path.with_suffix('.evidence.json').read_text())
            self.assertEqual(len(pages), expected['expected_pages'])
            self.assertGreater(expected['reviewed_note_pages'], 1)
            all_text = '\n'.join(text for page in pages for text in page)
            self.assertIn('FINALOBSERVATION42', all_text)
            self.assertIn('acceptance remains Not Evaluated.', all_text)
            self.assertEqual(evidence['metadata']['Observations'], notes['Observations'])
            self.assertTrue(all(row['acceptance'] == 'Not Evaluated'
                                for row in evidence['tables']['KPI Results']))

    def test_legacy_ordered_runs_keep_reference_and_candidate_identities(self):
        store, runs, compared, selected = self.fixture(count=1)
        legacy = [{key: value for key, value in run.items() if key != 'role'} for run in runs]
        original = [set(run) for run in legacy]
        tables, info = report_tables(store, legacy, compared, view={'cursors': [1., 4.]}, selected_ids=selected)
        self.assertEqual({row['run'] for row in info['Cursor interval evidence']}, {'reference', 'candidate'})
        self.assertEqual({row['Run'] for row in info['Source timing quality']}, {'reference', 'candidate'})
        for row in tables['Comparison Results'].to_dict('records'):
            self.assertEqual(row['baseline_evidence']['recording'], 'Reference recording')
            self.assertEqual(row['candidate_evidence']['recording'], 'Candidate recording')
        self.assertEqual([set(run) for run in legacy], original)
        self.assertFalse(report_preview(runs)['page_count_exact'])
        self.assertTrue(report_preview(runs, store=store)['page_count_exact'])

    def test_preview_resolves_projected_selection_like_export(self):
        store,runs,_,selected=self.fixture(count=1,comparison=False)
        source=store.channels[selected[0]]
        projected=store.add(Channel('Power · Analysis',source.x,source.y,unit=source.unit,
            coordinate_unit=source.coordinate_unit,dataset_id=source.dataset_id,source=source.source,
            parents=[source.id],attributes={'engineering_projection':True}))
        source_view={'cursors':[1.,4.],'cursor_locks':[{'channel':source.id,'index':2},None]}
        projection_view={'cursors':[1.,4.],'cursor_locks':[{'channel':projected.id,'index':2},None]}
        direct=report_preview(runs,view=source_view,selected_ids=[source.id],store=store)
        projection=report_preview(runs,view=projection_view,selected_ids=[projected.id],store=store)
        self.assertEqual(projection,direct)
        self.assertEqual(projection_view['cursor_locks'][0]['channel'],projected.id)

    def test_workbook_notes_and_source_labels_are_literal_while_results_stay_numeric(self):
        from openpyxl import load_workbook
        store,runs,_,selected=self.fixture(count=2,comparison=False)
        source=store.channels[selected[1]]
        source.name='=Literal source label'
        # Analyze the final source identity, as the real report workflow does.
        # A prior analysis deliberately retains its earlier mapping snapshot.
        runs=[analyze_run(store,runs[0]['profile'],dataset_id=source.dataset_id)]
        runs[0]['role']='analysis'
        metadata={'Observations':'=literal user note, not a formula',
                  'Prepared by':'=Literal author','Custom metadata':'=Literal custom value'}
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'Literal.xlsx'
            tables=export_engineering(path,store,runs,metadata=metadata,selected_ids=selected)
            book=load_workbook(path,data_only=False)
            self.assertEqual(book.sheetnames,SHEETS)
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        self.assertNotEqual(cell.data_type,'f',sheet.title+' contains a formula instead of report text')
            fields={row[0].value:row[1] for row in book['Test Metadata'].iter_rows(min_row=2)}
            for key,value in metadata.items():
                self.assertEqual(fields[key].value,value)
                self.assertEqual(fields[key].data_type,'s')
            self.assertTrue(any(cell.value==source.name and cell.data_type=='s'
                                for row in book['Channel Mapping'] for cell in row))
            numeric_sheet=book['KPI Results']
            value_column=[cell.value for cell in numeric_sheet[1]].index('value')+1
            expected=tables['KPI Results']['value'].tolist()
            known=0
            for row,number in enumerate(expected,2):
                if isinstance(number,(float,int,np.floating,np.integer)) and np.isfinite(number):
                    cell=numeric_sheet.cell(row,value_column)
                    self.assertEqual(cell.data_type,'n')
                    self.assertAlmostEqual(cell.value,number,places=12)
                    known+=1
            self.assertGreater(known,0)


if __name__ == '__main__':
    unittest.main()
