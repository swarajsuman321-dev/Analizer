"""Reusable channel presets, original-source exports, and recorded phase spans."""
import json
from pathlib import Path
import re
import numpy as np
import pandas as pd
from app_support import atomic_output,write_excel_frame

STYLES={'Continuous (-)':'solid','Dashed (--)':'dash','Dotted (:)':'dot','Dash-dot (-.)':'dashdot'}


def validate_preset(record):
    from PySide6.QtGui import QColor
    if not isinstance(record,dict):raise ValueError('Preset must be a JSON object')
    if 'version' in record and record['version'] not in (2,3):raise ValueError('Unsupported preset version')
    signals=record.get('signals',{}) if 'version' in record else record
    options=record.get('options',{}) if 'version' in record else {}
    if not isinstance(signals,dict) or not isinstance(options,dict):raise ValueError('Invalid preset signals/options')
    for key,entry in signals.items():
        if not isinstance(entry,dict):raise ValueError(f'Invalid channel settings: {key}')
        width=float(entry.get('width',1))
        if not np.isfinite(width) or not 0<width<=20:raise ValueError('Line width must be finite and between 0 and 20')
        for name in ('enabled','twin'):
            if name in entry and not isinstance(entry[name],bool):raise ValueError(f'{name} must be true or false')
        if not QColor(entry.get('color','#58b6ff')).isValid():raise ValueError('Invalid channel color')
        panel=entry.get('panel','Panel 1')
        if panel not in [f'Panel {i}' for i in range(1,11)]:raise ValueError('Choose panel 1 through 10')
        if entry.get('style','Continuous (-)') not in {*STYLES,*STYLES.values()}:raise ValueError('Invalid line style')
        bounds=[float(entry[name]) if str(entry.get(name,'')).strip() else None for name in ('ymin','ymax')]
        if any(v is not None and not np.isfinite(v) for v in bounds):raise ValueError('Axis bounds must be finite')
        if all(v is not None for v in bounds) and bounds[0]>=bounds[1]:raise ValueError('Y minimum must be less than maximum')
    for key in ('offset_minutes','font_size'):
        if key in options and not np.isfinite(float(options[key])):raise ValueError(f'{key} must be finite')
    if not 7<=float(options.get('font_size',9))<=20:raise ValueError('Font size must be between 7 and 20')
    for key in ('show_grid','show_legends','annotate_peak','highlight_overlap'):
        if key in options and not isinstance(options[key],bool):raise ValueError(f'{key} must be true or false')
    return signals,options


def channel_preset(store,workspace,options):
    signals={};visible={id for pane in workspace.visible_panes() for id in pane.ids}
    for c in store.channels.values():
        source=c.attributes.get('source_column',c.name);key=source
        if key in signals:key=f'{source} [{c.source or c.dataset_id}]'
        style=c.attributes.get('plot_style',{})
        signals[key]={'source_col':source,'source_path':c.source,'name':c.name,'unit':c.unit,'enabled':c.id in visible,
                      'panel':f'Panel {next((i+1 for i,p in enumerate(workspace.panes) if c.id in p.ids),1)}',
                      'twin':style.get('side','left')=='right','style':style.get('style','solid'),'width':style.get('width',1),
                      'color':style.get('color','#58b6ff'),'ymin':style.get('ymin',''),'ymax':style.get('ymax','')}
    return {'version':3,'signals':signals,'options':options}


def recorded_phases(channels):
    phase=next((c for c in channels if c.kind=='discrete' and any(word in (c.name+' '+c.attributes.get('signal_kind','')).lower() for word in ('stage','stwf','code','phase'))),None)
    if phase is None:return []
    x,y=phase.values();result=[]
    for i in range(max(0,len(x)-1)):
        value=y[i];label=f'{value:g}' if np.isfinite(value) else 'Unknown'
        if x[i+1]<=x[i]:continue
        if result and result[-1]['label']==label and result[-1]['end']==x[i]:result[-1]['end']=float(x[i+1])
        else:result.append({'start':float(x[i]),'end':float(x[i+1]),'label':label,'unit':phase.coordinate_unit})
    return result


def export_full_workbook(path,store,workspace=None,scorecard=None):
    from app_reporting import interval_tables
    from app_presentation import match_color
    from openpyxl.styles import PatternFill
    with atomic_output(path) as temporary:
        with pd.ExcelWriter(temporary,engine='openpyxl') as writer:
            used=set()
            def sheet(name,frame):
                clean=re.sub(r'[\\/*?:\[\]]','_',name)[:31] or 'Data';name=clean;number=2
                while name.casefold() in used:name=clean[:27]+'_'+str(number);number+=1
                used.add(name.casefold());write_excel_frame(writer,frame,name)
            if scorecard is not None:sheet('Comparison',scorecard)
            for dataset in store.datasets.values():sheet(dataset.frame.attrs.get('studio_sheet',dataset.name),dataset.frame)
            if workspace is not None and any(p['ids'] for p in workspace['panes']):
                data,summary=interval_tables(store,workspace);sheet('Active_Interval',data);sheet('Cursor_Summary',summary)
            if not used:raise ValueError('Load data before exporting a workbook')
            for ws in writer.book:
                ws.freeze_panes='A2';headers=[cell.value for cell in ws[1]]
                if 'match_percent' in headers:
                    col=headers.index('match_percent')+1
                    for row in range(2,ws.max_row+1):
                        cell=ws.cell(row,col);color=match_color(cell.value)
                        if color:cell.fill=PatternFill('solid',fgColor=color[1:]);cell.number_format='0.0"%"'
