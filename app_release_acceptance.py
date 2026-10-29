"""Executable acceptance checks shared by source and all Windows editions.

The application itself imports this module for ``--acceptance-check``.  A test
runner cannot substitute a Python stub for the frozen workbench: the result
records the executable, frozen status, edition and actual Qt platform.  Physical
DPI, graphics drivers, accessibility and manual mouse checks remain explicitly
unrun; offscreen geometry is not evidence for those claims.
"""
from contextlib import contextmanager
from datetime import datetime, timezone
import faulthandler
import hashlib
from itertools import combinations
import json
import os
from pathlib import Path
import sys
import time
import traceback
import uuid


MANUAL_WINDOWS_CHECKS = [
    'Physical wheel, box selection and keyboard navigation in every zoom mode',
    'Windows 100%, 125%, 150% and 200% display scaling',
    'Moving between monitors with different scaling and graphics drivers',
    'Keyboard-only task completion and screen-reader announcements',
    'Long-running acquisition-sized workloads and antivirus policy interaction',
]

# This finite contract inventory prevents a green but accidentally shortened
# runner from satisfying a release gate simply by retaining its category labels.
# New checks may be added without weakening any of the already approved ones.
REQUIRED_CHECK_IDS = frozenset({
    'REAL-FIXTURE-INTEGRITY-IMPORT', 'ALL-15-FILE-COMBINATIONS',
    'PER-FILE-CHECKBOX-AND-CLOSE-DATASET', 'TDMS-FIRST-REVIEW-REUSE-UNKNOWN-ONLY',
    'SYNTHETIC-SOURCE-EDGE-CASES', 'LOCKED-SAMPLE-MARKER-HUD-EXPORT-RESTORE',
    'GRAPH-HEADER-CONTEXT-SAFE-NATIVE-FFT', 'NON-TEMPORAL-AND-UNLOCKED-INDEPENDENCE',
    'NUMERICAL-UNITS-AND-SAFETY-REJECTIONS',
    'PROCESSING-PREFLIGHT-VALID-DEFAULT-AND-EXPLANATIONS',
    'QT-WHEEL-BOX-KEYBOARD-SHARED-MODES',
    'EXACT-REAL-SOURCE-WORKSPACE-ROUNDTRIP',
    'REAL-INTERVAL-EXPORT-FORMATS-AND-AGREEMENT', 'SOURCE-FILES-UNCHANGED',
    'DEFAULT-PANEL-FOLDS-EXPAND-PRESERVE-READOUT',
    'REAL-TRACE-COLOR-SCOPE-PRESET-EXPORT',
    'REAL-LOADED-DATA-COMPARISON-ISOLATION',
    'REAL-SELECTED-GRAPH-FULLSCREEN-VALUES-RESTORE',
    'PERSISTENT-DIALOG-FOLDERS-RESTART-CANCEL-SCOPES',
    'REAL-WORKSPACE-FULLSCREEN-SPLITS-HOVER-RESTORE',
    'REAL-TWO-SIGNAL-COMPARISON-EVIDENCE',
} | {f'LAYOUT-{layout}-{theme}-{width}'
     for layout in ('Single', 'Horizontal split', 'Vertical split', '2×2', 'Stacked', '3 stacked', '10 panels')
     for theme in ('Dark', 'Light') for width in (1024, 1536)})


def sha256_file(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def viewport_time_x(pane, value):
    """Map a timestamp through the complete floating viewport transform.

    Each plot owns its own scene.  Comparing scene coordinates, or rounding with
    mapFromScene(), can miss different canvas gutters at the same numeric time.
    Qt returns logical pixels here; this is not physical-DPI certification.
    """
    from PySide6 import QtCore
    point = pane.getViewBox().mapViewToScene(QtCore.QPointF(float(value), 0))
    viewport = pane.viewportTransform().map(point)
    return float(pane.viewport().mapToGlobal(QtCore.QPoint(0, 0)).x()+viewport.x())


def time_pixel_alignment(panes, columns, *, tolerance=.5):
    """Assert equal numeric time bounds also draw at equal X pixels per column."""
    import numpy as np
    assert columns > 0
    groups = {}
    for index, pane in enumerate(panes):
        if pane.ids and pane.coordinate and pane.coordinate[0] == 'time':
            groups.setdefault((index % columns, tuple(pane.coordinate)), []).append((index, pane))
    observed = []
    for (column, coordinate), members in groups.items():
        first = members[0][1]
        bounds = np.asarray(first.getViewBox().viewRange()[0], dtype=float)
        timestamps = [float(bounds[0]+(bounds[1]-bounds[0])*fraction) for fraction in (0, .1, .5, .9, 1)]
        records = []
        for index, pane in members:
            np.testing.assert_allclose(pane.getViewBox().viewRange()[0], bounds, rtol=0, atol=1e-7,
                err_msg='Pixel alignment must preserve identical numeric time limits')
            positions = [viewport_time_x(pane, timestamp) for timestamp in timestamps]
            assert np.isfinite(positions).all() and positions[-1] > positions[0], 'Invalid time-to-viewport transform'
            records.append({'pane_index': index, 'time_pixel_x': positions,
                            'axis_units': list(dict.fromkeys(pane.bindings.values()))})
        pixels = np.asarray([record['time_pixel_x'] for record in records])
        spreads = np.ptp(pixels, axis=0)
        assert np.all(spreads <= tolerance+1e-7), (
            f'Common timestamps differ by {spreads.tolist()} logical pixels in grid column {column}; '
            f'numeric bounds agree but the drawing gutters do not (limit {tolerance})')
        observed.append({'column': column, 'coordinate': list(coordinate), 'numeric_bounds': bounds.tolist(),
                         'timestamps': timestamps, 'maximum_pixel_spread': float(spreads.max()),
                         'tolerance_logical_pixels': tolerance, 'panes': records})
    assert observed, 'No temporal pane was available for pixel-alignment acceptance'
    return observed


def verify_comparison_interval_metrics(result):
    """Independently integrate exported linear differences, not engine totals.

    Simpson integration is exact for a linear difference and its square. For
    absolute error, split an interval at a sign crossing before integrating.
    This cross-check consumes the complete exported evidence and also rejects
    omitted/overlapping timeline intervals or unsupported values marked known.
    """
    import math
    import numpy as np
    assert result.get('status') == 'compared'
    scope, metrics = result['scope'], result['metrics']
    rows = result['differences']
    assert rows and scope['span_min'] > 0
    end = float(scope['start_min'])
    duration = bias_area = square_area = absolute_area = 0.
    maximum = 0.
    supported = 0
    for row in rows:
        left, right = float(row['start_min']), float(row['end_min'])
        assert np.isclose(left, end, rtol=0, atol=1e-7), 'Difference evidence omits or overlaps a timeline interval'
        h = right-left
        assert h > 0 and np.isclose(h, row['duration_min'], rtol=1e-9, atol=1e-9)
        end = right
        if not row['supported']:
            assert row['delta_start'] is None and row['delta_end'] is None
            continue
        a, b = float(row['delta_start']), float(row['delta_end'])
        assert math.isfinite(a) and math.isfinite(b)
        midpoint = (a+b)/2
        duration += h
        bias_area += h*(a+4*midpoint+b)/6
        square_area += h*(a*a+4*midpoint*midpoint+b*b)/6
        fractions = [0., 1.]
        if a*b < 0:
            fractions.insert(1, -a/(b-a))
        for x, y in zip(fractions, fractions[1:]):
            absolute_area += h*(y-x)*(abs(a+(b-a)*x)+abs(a+(b-a)*y))/2
        maximum = max(maximum, abs(a), abs(b))
        supported += 1
    assert np.isclose(end, scope['end_min'], rtol=0, atol=1e-7)
    assert duration > 0 and supported > 0
    expected = {'compared_duration_min': duration, 'coverage_fraction': duration/scope['span_min'],
                'mean_bias': bias_area/duration, 'mean_absolute_difference': absolute_area/duration,
                'rmse': math.sqrt(square_area/duration), 'max_absolute_difference': maximum}
    for field, value in expected.items():
        assert np.isclose(metrics[field], value, rtol=1e-9, atol=1e-8), 'Exported intervals disagree with '+field
    return {'supported_intervals': supported, 'complete_intervals': len(rows), **expected}


@contextmanager
def isolated_environment(output):
    """Keep diagnostics away from a user's templates, outputs and settings."""
    overrides = {'ANALIZER_SETTINGS_DIR': str(output/'settings'),
                 'ANALIZER_OUTPUT_DIR': str(output/'outputs'),
                 'MPLCONFIGDIR': str(output/'matplotlib'),
                 'XDG_CACHE_HOME': str(output/'cache'),
                 'QT_QPA_PLATFORM': os.environ.get('QT_QPA_PLATFORM', 'offscreen')}
    previous = {key: os.environ.get(key) for key in overrides}
    os.environ.update(overrides)
    try:
        yield
    finally:
        for key, value in previous.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value


def acceptance_worker_events(app):
    """Deliver real GUI signals while yielding Python to its background worker.

    QTest.qWait can retain the GIL through much of its native polling loop.
    Using it repeatedly around a Python report writer can serialize compression
    and make the diagnostic workload much slower than the normal app event loop.
    Python sleep releases the GIL; GUI callbacks still run on the GUI thread.
    """
    app.processEvents()
    time.sleep(.025)
    app.processEvents()


def wait_for_acceptance_workers(window, events, *, timeout, cancel_timeout=30):
    """Keep the original deadline a failure, then drain queued cancellation.

    A failed asynchronous check must not leave its worker running after its
    nonblocking-message boundary is restored. Deliver the normal Cancel action
    and process its queued completion for a bounded cleanup period. Successful
    cleanup never turns the missed original deadline into a pass.
    """
    deadline = time.monotonic()+timeout
    while (window.busy or window.workers) and time.monotonic() < deadline:
        events()
    if not window.busy and not window.workers:
        return
    window.cancel_operation()
    cleanup_deadline = time.monotonic()+cancel_timeout
    while (window.busy or window.workers) and time.monotonic() < cleanup_deadline:
        events()
    state = ('cancellation completed' if not window.busy and not window.workers
             else 'cancellation is still pending')
    raise AssertionError('The comparison worker did not finish within six minutes; '+state)


class AcceptanceRecorder:
    """Persist completed checks even when a later application callback fails."""
    def __init__(self, path, metadata=None):
        self.path = Path(path)
        self.result = {**(metadata or {}), 'checks': [], 'callback_errors': [],
                       'runner_errors': [], 'expected_safety_rejections': [],
                       'manual_checks_unrun': list(MANUAL_WINDOWS_CHECKS)}
        self.current = 'startup'

    def save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(self.result, indent=2, ensure_ascii=False,
                                       default=str)+'\n', encoding='utf-8')

    def check(self, name, operation, category, requires=()):
        self.current = name
        self.result['current_check'] = name
        self.result['check_started_at_utc'] = datetime.now(timezone.utc).isoformat()
        self.result['current_stage'] = 'Starting check'
        self.save()
        started = time.perf_counter()
        complete = {row['id']: row['status'] for row in self.result['checks']}
        blockers = [dependency for dependency in requires if complete.get(dependency) != 'passed']
        if blockers:
            row = {'id': name, 'category': category, 'status': 'unrun',
                   'blocked_by': blockers, 'seconds': 0}
            self.result['checks'].append(row)
            self.save()
            return row
        try:
            observed = operation()
            row = {'id': name, 'category': category, 'status': 'passed',
                   'observed': observed}
        except AssertionError:
            row = {'id': name, 'category': category, 'status': 'unmet',
                   'error': traceback.format_exc()}
        except Exception:
            error = {'id': name, 'category': category, 'error': traceback.format_exc()}
            self.result['runner_errors'].append(error)
            row = {'id': name, 'category': category, 'status': 'execution_error',
                   'error': error['error']}
        row['seconds'] = round(time.perf_counter()-started, 3)
        self.result['checks'].append(row)
        self.save()
        return row

    def stage(self, message):
        """Retain the last real UI progress message if a frozen process stalls."""
        if message != self.result.get('current_stage'):
            self.result['current_stage'] = str(message)
            self.result['stage_updated_at_utc'] = datetime.now(timezone.utc).isoformat()
            self.save()

    def finish(self, required_checks=()):
        present = {row['id'] for row in self.result['checks']}
        for missing in sorted(set(required_checks)-present):
            self.result['checks'].append({'id': missing, 'category': 'coverage',
                'status': 'unrun', 'seconds': 0,
                'reason': 'The required acceptance contract did not execute'})
        rows = self.result['checks']
        self.result['summary'] = {
            'checks': len(rows), 'passed': sum(r['status'] == 'passed' for r in rows),
            'unmet': sum(r['status'] == 'unmet' for r in rows),
            'unrun': sum(r['status'] == 'unrun' for r in rows),
            'runner_errors': len(self.result['runner_errors']),
            'application_callback_errors': len(self.result['callback_errors']),
            'expected_safety_rejections': len(self.result['expected_safety_rejections']),
            'manual_unrun': len(self.result['manual_checks_unrun'])}
        self.result['ok'] = bool(rows) and all(r['status'] == 'passed' for r in rows) and not (
            self.result['callback_errors'] or self.result['runner_errors'])
        self.result['finished_at_utc'] = datetime.now(timezone.utc).isoformat()
        self.save()
        return self.result


@contextmanager
def acceptance_stack_diagnostics(output):
    """Capture all Python threads even when a windowed EXE has no stderr.

    This watchdog observes the diagnostic CLI only. It does not interrupt work,
    extend its deadline, or turn an unfinished check into a successful one.
    """
    with (Path(output)/'acceptance-stacks.txt').open('w', encoding='utf-8') as stream:
        faulthandler.dump_traceback_later(120, repeat=True, file=stream)
        try:
            yield
        finally:
            faulthandler.cancel_dump_traceback_later()


def run_acceptance(samples, output):
    """Exercise the real workbench; return a JSON-compatible finite test report.

    ``samples`` is an external directory containing the four committed recording
    fixtures and manifest.  Test recordings are deliberately not bundled into
    application distributions.  Missing or changed fixtures fail the gate.
    ``output`` must be a disposable diagnostic directory, never a user workspace.
    """
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    # No display server is necessary on the native CI runner.  A caller can
    # deliberately select a real platform plugin for separate visual checks.
    from app_version import VERSION
    from app_installation import installed_root
    frozen = bool(getattr(sys, 'frozen', False))
    installed = bool(installed_root())
    distribution_form = ('source' if not frozen else 'installed' if installed else
                         'standalone-folder' if (Path(sys.executable).parent/'_internal').is_dir()
                         else 'standalone-onefile')
    recorder = AcceptanceRecorder(output/'acceptance.json', {
        'schema': 1, 'version': VERSION, 'platform': sys.platform,
        'run_id': os.environ.get('ANALIZER_ACCEPTANCE_RUN_ID') or str(uuid.uuid4()),
        'started_at_utc': datetime.now(timezone.utc).isoformat(),
        'frozen': frozen, 'executable': str(Path(sys.executable).resolve()),
        'edition': 'source' if not frozen else 'installed' if installed else 'standalone',
        'distribution_form': distribution_form,
        'qt_platform_requested': os.environ.get('QT_QPA_PLATFORM', 'offscreen'), 'inputs': []})
    with isolated_environment(output), acceptance_stack_diagnostics(output):
        _run_workbench(recorder, Path(samples).resolve() if samples else None, output)
    return recorder.finish(required_checks=REQUIRED_CHECK_IDS)


def _run_workbench(recorder, samples, output):
    from copy import deepcopy
    import numpy as np
    import pandas as pd
    from PySide6 import QtCore, QtGui, QtWidgets, QtTest
    from app_channel_store import Channel, ChannelStore, Dataset
    from app_ingestion import ingest
    from app_plot_workspace import LAYOUTS
    from app_reporting import export_report, interval_tables
    from app_signal_processing import strict_clock, transform, formula
    from app_tdms_templates import ParameterTemplates, ParameterReview, identity
    from app_workbench import AdvancedStudioWindow
    from app_workspace_io import save_workspace, load_workspace

    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    recorder.result['qt_platform'] = app.platformName()
    old_hook = sys.excepthook
    def callback_error(kind, value, tb):
        recorder.result['callback_errors'].append({'scenario': recorder.current,
            'error': ''.join(traceback.format_exception(kind, value, tb))})
        recorder.save()
    sys.excepthook = callback_error
    window = None
    def events():
        app.processEvents()
        QtTest.QTest.qWait(25)
        app.processEvents()
    try:
        window = AdvancedStudioWindow()
        window.resize(1536, 960)
        window.show()
        window.statusBar().messageChanged.connect(recorder.stage)
        events()
        store = window.store
        workspace = window.dashboard
        fixture_stores = {}
        manifest = None

        def load_fixtures():
            nonlocal manifest
            assert samples and (samples/'manifest.json').is_file(), 'The real-fixture manifest is required'
            manifest = json.loads((samples/'manifest.json').read_text(encoding='utf-8'))
            assert len(manifest['files']) == 4, 'Expected the four approved recording fixtures'
            for entry in manifest['files']:
                path = samples/entry['filename']
                assert path.is_file() and path.stat().st_size == entry['bytes'], 'Fixture size mismatch: '+path.name
                assert sha256_file(path) == entry['sha256'], 'Fixture digest mismatch: '+path.name
                other = ChannelStore()
                ingest(other, path)
                assert other.datasets and other.channels, 'No source data imported: '+path.name
                fixture_stores[entry['role']] = other
                recorder.result['inputs'].append({**entry, 'channels': len(other.channels),
                    'source_rows': [len(dataset.frame) for dataset in other.datasets.values()]})
            store.datasets.update({key: value for other in fixture_stores.values() for key, value in other.datasets.items()})
            with store.batch():
                for other in fixture_stores.values():
                    for channel in other.channels.values():
                        store.add(channel)
            events()
            return {'files': len(fixture_stores), 'channels': len(store.channels)}
        recorder.check('REAL-FIXTURE-INTEGRITY-IMPORT', load_fixtures, 'real_recordings')

        def real_channel(role, unit=None):
            candidates = [store.channels[channel_id] for channel_id in fixture_stores[role].channels
                          if store.channels[channel_id].coordinate_kind == 'time'
                          and store.channels[channel_id].coordinate_unit == 'min'
                          and np.isfinite(store.channels[channel_id].values()[1]).any()]
            if unit is not None:
                candidates = [channel for channel in candidates if channel.unit == unit]
            assert candidates, f'No usable {role} source channel with declared unit {unit!r}'
            return candidates[0]

        def folded_panels_preserve_readout():
            controls = window.panel_folding
            folded = ('Properties', 'Channel Math', 'Channel Table', 'Interval Measurements')
            assert all(controls.is_folded(name) for name in folded), 'The four approved default folds were not applied'
            assert all(not controls.is_folded(name) for name in ('Engineering Controls', 'Readout', 'Data Portal'))
            assert window.docks['Engineering Controls'].widget().isVisible(), 'Engineering Controls is not the default available tab'
            source = real_channel('tdms', '°C')
            workspace.active.bind([source.id])
            window._readouts()
            events()
            original_widgets = {name: dock.widget() for name, dock in window.docks.items()}
            before = deepcopy(workspace.state())
            before_readout = [[window.readout.item(row, column).text()
                               for column in range(window.readout.columnCount())]
                              for row in range(window.readout.rowCount())]
            assert before_readout, 'The retained Readout has no actual source values'
            assert controls.button.accessibleName() and controls.button.menu() is controls.menu
            toolbar = window.main_toolbar
            extension = toolbar.findChild(QtWidgets.QToolButton, 'qt_toolbar_ext_button')
            assert controls.button.isVisible() or (extension is not None and extension.isVisible()), 'Panels has no visible button or overflow access'
            for name in folded:
                action = controls.actions[name]
                assert action in controls.menu.actions() and action.isEnabled() and not action.icon().isNull()
                assert action.text() == 'Expand '+name and not action.isChecked()
                # Exercise the same QAction available from the visible Panels menu.
                action.trigger()
                events()
                assert not controls.is_folded(name) and action.isChecked()
                assert action.text() == 'Fold '+name
                assert window.docks[name].widget() is original_widgets[name], 'Expanding replaced panel contents'
                controls.set_folded(name, True)
                events()
            controls.apply_defaults(folded)
            events()
            assert workspace.state() == before, 'Folding changed the plotted channels, cursors or view'
            after_readout = [[window.readout.item(row, column).text()
                              for column in range(window.readout.columnCount())]
                             for row in range(window.readout.rowCount())]
            assert after_readout == before_readout, 'Panel folding discarded the source readout'
            assert all(window.docks[name].widget() is widget for name, widget in original_widgets.items())
            assert not controls.is_folded('Engineering Controls') and not controls.is_folded('Readout')
            assert window.docks['Readout'].widget().isVisible(), 'Default folds hid the retained Readout'
            return {'default_folded': list(folded), 'source_role': 'tdms',
                    'expanded_and_refolded': len(folded), 'preserved_readout_rows': len(before_readout),
                    'engineering_and_readout_available': True}
        recorder.check('DEFAULT-PANEL-FOLDS-EXPAND-PRESERVE-READOUT', folded_panels_preserve_readout,
                       'panel_folding', requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        def persistent_dialog_folders():
            import importlib
            from unittest.mock import patch
            import app_dialog_history as history
            from app_lifecycle import settings_directory
            from app_output import output_path

            # This is the actual source/frozen workbench route. Mock only the
            # native dialog boundary so no operator interaction or data import
            # can occur while testing its persisted initial folder.
            window.tabs.setCurrentWidget(workspace)
            workspace.active.bind([real_channel('tdms', '°C').id])
            events()
            views = [deepcopy(view.state()) for view in window.workspaces]
            channels, datasets = dict(store.channels), dict(store.datasets)
            workers = list(window.workers)
            history_path = settings_directory()/history.HISTORY_FILENAME
            previous = history_path.read_bytes() if history_path.exists() else None
            folders = {purpose: output/'dialog-folder-check'/name for purpose, name in (
                ('OpenData', "José's 測定 données"), ('ExportReport', 'Report exports'),
                ('WorkspaceSave', 'Workspace saves'), ('WorkspaceOpen', 'Workspace imports'))}
            for folder in folders.values():
                folder.mkdir(parents=True, exist_ok=True)
            cancelled_routes = []

            def cancel_route(method, operation, expected, response):
                before = history_path.read_bytes()
                with patch.object(QtWidgets.QFileDialog, method, return_value=response) as dialog:
                    result = operation()
                dialog.assert_called_once()
                assert dialog.call_args.args[0] is window
                assert dialog.call_args.args[2] == str(expected), 'The dialog ignored its persisted purpose folder or configured first-save default'
                assert history_path.read_bytes() == before, 'Cancelling a dialog rewrote personal history'
                assert list(window.workers) == workers and not window.busy, 'Cancelling a dialog started a data worker'
                cancelled_routes.append(method)
                return result

            try:
                history_path.unlink(missing_ok=True)
                assert history.remember_path('OpenData', folders['OpenData']/'previous recording.csv')
                history = importlib.reload(history)
                assert history.start_directory('OpenData') == str(folders['OpenData'])
                assert 'previous recording.csv' not in history_path.read_text(encoding='utf-8'), 'History retained a selected measurement filename'
                assert cancel_route('getOpenFileNames', window.open_data, folders['OpenData'], ([], '')) is False
                # A first export/save uses the configured output area even
                # after a different input recording folder was remembered.
                cancel_route('getSaveFileName', window.report, output_path('Signal report.pdf'), ('', ''))
                cancel_route('getSaveFileName', window.save, output_path('Workspace.json', 'workspaces'), ('', ''))
                for purpose in ('ExportReport', 'WorkspaceSave', 'WorkspaceOpen'):
                    assert history.remember_path(purpose, folders[purpose]/'previous selection.json')
                history = importlib.reload(history)
                cancel_route('getOpenFileNames', window.open_data, folders['OpenData'], ([], ''))
                cancel_route('getSaveFileName', window.report, folders['ExportReport']/'Signal report.pdf', ('', ''))
                cancel_route('getSaveFileName', window.save, folders['WorkspaceSave']/'Workspace.json', ('', ''))
                cancel_route('getOpenFileName', window.restore, folders['WorkspaceOpen'], ('', ''))
                assert [view.state() for view in window.workspaces] == views, 'Cancelled dialogs changed a plotted workspace'
                assert set(store.channels) == set(channels) and set(store.datasets) == set(datasets)
                assert all(store.channels[key] is value for key, value in channels.items()), 'Cancelled dialogs changed source channels'
                assert all(store.datasets[key] is value for key, value in datasets.items()), 'Cancelled dialogs changed source datasets'
                return {'persisted_across_module_reload': True, 'unicode_folder': True,
                        'purposes': list(folders), 'cancelled_routes': len(cancelled_routes),
                        'configured_first_save_fallbacks': 2, 'history_and_sources_unchanged_on_cancel': True}
            finally:
                if previous is None:
                    history_path.unlink(missing_ok=True)
                else:
                    history_path.write_bytes(previous)
        recorder.check('PERSISTENT-DIALOG-FOLDERS-RESTART-CANCEL-SCOPES', persistent_dialog_folders,
                       'dialog_history', requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        def direct_real_trace_colors():
            from app_studio_state import channel_preset
            from app_workspace_io import render_report
            first, other = real_channel('tdms', '°C'), real_channel('machine', 'Pa')
            window.tabs.setCurrentWidget(workspace)
            workspace.set_layout('Vertical split')
            workspace.set_time_lock(False)
            workspace.panes[0].bind([first.id])
            workspace.panes[1].bind([other.id])
            workspace.set_active(workspace.panes[0])
            window.comparison.active.bind([first.id, other.id])
            workspace.panes[0].lock_cursor(0, first.id)
            workspace.panes[0].remember_view()
            events()
            panes = [pane for view in window.workspaces for pane in view.panes]
            before = [(pane, deepcopy(pane.state()), deepcopy(pane.history), pane.history_index)
                      for pane in panes]
            views = [(view, view.layout_name, view.active) for view in window.workspaces]
            # ChannelStore's normal metadata update replaces its immutable
            # Channel value object. Require source identity and exact readings,
            # rather than incidental NumPy object identity across revisions.
            source_fields = ('id', 'name', 'unit', 'kind', 'coordinate_kind',
                             'coordinate_unit', 'dataset_id', 'source', 'shift',
                             'masks', 'parents', 'recipe')
            originals = {channel.id: {
                'identity': {name: deepcopy(getattr(channel, name)) for name in source_fields},
                'x': channel.x.copy(), 'y': channel.y.copy(),
                'raw_y': None if channel.raw_y is None else channel.raw_y.copy(),
            } for channel in (first, other)}
            source_frames = {channel.dataset_id: store.datasets[channel.dataset_id].frame.copy(deep=True)
                             for channel in (first, other)}
            backgrounds = [pane.backgroundBrush().color().rgba() for pane in panes]
            unchanged_color = workspace.panes[1].curves[other.id][1].pen.color().name()
            color = '#c94b8f'
            assert window.trace_colors.set_color([first.id], color) == [first.id]
            events()
            assert [(view.layout_name, view.active) for view, _, _ in views] == [(layout, active) for _, layout, active in views]
            for pane, state, history, index in before:
                assert pane.state() == state and pane.history == history and pane.history_index == index, 'A trace-color edit reset navigation or cursor identity'
            assert [pane.backgroundBrush().color().rgba() for pane in panes] == backgrounds, 'Trace color changed graph backgrounds'
            for pane in (workspace.panes[0], window.comparison.active):
                assert pane.curves[first.id][1].pen.color().name() == color
                sample, _label = pane.legend.items[pane.ids.index(first.id)]
                assert sample.item.opts['pen'].color().name() == color, 'The legend swatch disagrees with its trace'
            assert workspace.panes[1].curves[other.id][1].pen.color().name() == unchanged_color
            for channel_id, original in originals.items():
                channel = store.channels[channel_id]
                assert {name: getattr(channel, name) for name in source_fields} == original['identity']
                np.testing.assert_array_equal(channel.x, original['x'])
                np.testing.assert_array_equal(channel.y, original['y'])
                if original['raw_y'] is None:
                    assert channel.raw_y is None
                else:
                    np.testing.assert_array_equal(channel.raw_y, original['raw_y'])
            for dataset_id, original in source_frames.items():
                pd.testing.assert_frame_equal(store.datasets[dataset_id].frame, original)
            assert store.channels[first.id].attributes['plot_style']['color'] == color
            preset = channel_preset(store, workspace, window.studio.options)
            original_column = first.attributes.get('source_column', first.name)
            assert any(entry['source_col'] == original_column and entry['source_path'] == first.source
                       and entry['color'] == color for entry in preset['signals'].values()), 'A saved preset dropped the direct trace color'
            path = output/'direct-real-trace-color.svg'
            render_report(path, store, workspace.state())
            assert color in path.read_text(encoding='utf-8').lower(), 'The exported plot omitted the chosen trace color'
            return {'source_roles': ['tdms', 'machine'], 'changed_channel': first.name,
                    'color': color, 'background_unchanged': True, 'scopes_checked': ['Analysis', 'Comparison', 'legend', 'preset', 'SVG']}
        recorder.check('REAL-TRACE-COLOR-SCOPE-PRESET-EXPORT', direct_real_trace_colors, 'trace_appearance',
                       requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        def arbitrary_loaded_comparison():
            comparison = window.comparison
            comparison.set_layout('Single')
            for pane in comparison.panes:
                pane.bind([])
            representatives = [real_channel(role) for role in fixture_stores]
            ids = [channel.id for channel in representatives]
            originals = {channel.id: (channel, channel.x, channel.y, channel.x.copy(), channel.y.copy())
                         for channel in representatives}
            before_analysis = deepcopy(workspace.state())
            source_ids, dataset_ids = set(store.channels), set(store.datasets)
            first = window.data_comparison.add_channels(ids[:2])
            events()
            pane = comparison.active
            anchor = representatives[0]
            pane.lock_cursor(0, anchor.id)
            pane.annotate('Marker', pane.cursor_lines[0].value(), text='Acceptance comparison annotation')
            events()
            retained = deepcopy(pane.state())
            second = window.data_comparison.add_channels(ids[2:])
            events()
            assert first['added'] == 2 and second['added'] == 2
            plotted = [channel_id for graph in comparison.visible_panes() for channel_id in graph.ids]
            assert len(plotted) == len(ids) and set(plotted) == set(ids), 'Compare Data copied or omitted source identities'
            assert len({store.channels[channel_id].dataset_id for channel_id in plotted}) == 4
            assert window.tabs.currentWidget() is comparison
            assert pane.cursor_locks == retained['cursor_locks'] and pane.annotations == retained['annotations']
            np.testing.assert_allclose(pane.getViewBox().viewRange(), retained['range'], rtol=0, atol=1e-7)
            assert workspace.state() == before_analysis, 'Adding loaded Comparison channels changed Analysis'
            repeated = window.data_comparison.add_channels(ids+ids)
            events()
            assert repeated == {'added': 0, 'already_present': 4, 'total': 4, 'datasets': 4}
            assert [channel_id for graph in comparison.visible_panes() for channel_id in graph.ids] == plotted
            assert set(store.channels) == source_ids and set(store.datasets) == dataset_ids
            for channel_id, (channel, x, y, expected_x, expected_y) in originals.items():
                assert store.channels[channel_id] is channel and channel.x is x and channel.y is y
                np.testing.assert_array_equal(channel.x, expected_x)
                np.testing.assert_array_equal(channel.y, expected_y)
            return {'roles': list(fixture_stores), 'original_channels': len(ids), 'datasets': 4,
                    'repeat_added': repeated['added'], 'analysis_unchanged': True,
                    'existing_comparison_lock_and_annotation_retained': True}
        recorder.check('REAL-LOADED-DATA-COMPARISON-ISOLATION', arbitrary_loaded_comparison, 'loaded_comparison',
                       requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        def selected_graph_fullscreen():
            from app_plot_workspace import PlotPane
            from app_presentation import CURSOR_COLORS
            from shiboken6 import isValid
            window.tabs.setCurrentWidget(workspace)
            workspace.set_layout('Vertical split')
            workspace.set_time_lock(True)
            temperature, pressure = real_channel('tdms', '°C'), real_channel('machine', 'Pa')
            first, selected = workspace.visible_panes()
            first.bind([temperature.id])
            selected.bind([pressure.id, temperature.id])
            workspace.set_active(selected)
            x, y = pressure.values()
            finite = np.isfinite(y)
            low, high = float(x[0]), float(x[-1])
            changes = np.flatnonzero(finite[:-1] & finite[1:] & (y[:-1] != y[1:])
                & (x[:-1] > low+.2*(high-low)) & (x[1:] < low+.6*(high-low)))
            assert len(changes), 'The real pressure fixture needs an interior changing source sample'
            index = int(changes[len(changes)//2])
            later = np.flatnonzero(finite & (x > x[index]+.05*(high-low)))
            assert len(later), 'The fixture needs a later finite Cursor B sample'
            last = int(later[len(later)//2])
            for cursor, sample in enumerate((index, last)):
                selected.lock_cursor(cursor, pressure.id)
                selected.cursor_locks[cursor]['index'] = sample
            selected.restore_ranges(selected.state())
            selected.setXRange(low, high, padding=0)
            events()
            parent = selected.parentWidget()
            slot = workspace.grid.getItemPosition(workspace.grid.indexOf(selected))
            panes = list(workspace.panes)
            curves = {channel_id: curve for channel_id, (_view, curve) in selected.curves.items()}
            dock_state = QtCore.QByteArray(window.saveState())
            geometry = QtCore.QRect(window.geometry())
            bindings = dict(selected.bindings)
            controller = window.graph_fullscreen
            try:
                host = controller.enter(selected)
                events()
                QtTest.QTest.qWait(50)
                app.processEvents()
                assert controller.is_active and host.isFullScreen() and host.isVisible()
                assert controller.pane is selected and controller.workspace is workspace
                assert host.findChildren(PlotPane) == [selected], 'Fullscreen expanded more than the selected graph'
                assert controller.readout.isVisible() and all(card.isVisible() for card in controller.cursor_cards)
                def verify_live_values():
                    # Inspect the UI after its ordinary queued updates. Calling
                    # refresh_readouts here would conceal a broken live signal.
                    rows = {controller.readout.item(row, 0).data(QtCore.Qt.ItemDataRole.UserRole): row
                            for row in range(controller.readout.rowCount())}
                    assert set(rows) == set(selected.ids)
                    row = rows[pressure.id]
                    for cursor, column in ((0, 2), (1, 3)):
                        source_index = selected.cursor_locks[cursor]['index']
                        expected = y[source_index]
                        assert selected.cursor_value(pressure.id, cursor) == expected
                        assert controller.readout.item(row, column).text() == f'{expected:.6g}', 'Fullscreen A/B values differ from the locked source samples'
                        assert f'{expected:.6g} Pa' in controller.cursor_cards[cursor].toolTip()
                        assert controller.readout.item(row, column).foreground().color().name() == CURSOR_COLORS[cursor].lower()
                        assert CURSOR_COLORS[cursor].lower() in controller.cursor_cards[cursor].styleSheet().lower()
                    assert controller.readout.item(row, 1).text() == 'Pa'
                    return row
                verify_live_values()
                for command in ('x', 'y', 'xy', 'pan', 'in', 'out', 'reset'):
                    action = workspace.navigation_actions[command]
                    assert action in controller.navigation_bar.actions() and action.isEnabled()
                    button = controller.navigation_bar.widgetForAction(action)
                    extension = controller.navigation_bar.findChild(QtWidgets.QToolButton, 'qt_toolbar_ext_button')
                    assert button.isVisible() or (extension is not None and extension.isVisible()), 'Fullscreen navigation has no visible button or overflow route'
                workspace.navigation_actions['x'].trigger()
                before_x = np.asarray(selected.getViewBox().viewRange()[0])
                before_y = [list(view.viewRange()[1]) for view in selected.y_views()]
                workspace.navigation_actions['in'].trigger()
                events()
                zoomed = list(selected.getViewBox().viewRange()[0])
                assert np.isclose(np.ptp(zoomed)/np.ptp(before_x), .8), 'Fullscreen zoom did not operate on the selected graph'
                for view, bounds in zip(selected.y_views(), before_y):
                    np.testing.assert_allclose(view.viewRange()[1], bounds, rtol=0, atol=1e-7)
                selected.step_cursor(0, 1)
                events()
                assert selected.cursor_locks[0]['index'] == index+1
                verify_live_values()
                source_positions = [line.value() for line in selected.cursor_lines]
            finally:
                controller.exit()
                events()
            assert not controller.is_active and not controller.fullscreen_window.isVisible()
            assert workspace.panes == panes and selected.parentWidget() is parent
            assert workspace.grid.getItemPosition(workspace.grid.indexOf(selected)) == slot
            assert selected.bindings == bindings and all(selected.curves[channel_id][1] is curve for channel_id, curve in curves.items())
            assert window.saveState() == dock_state and window.geometry() == geometry, 'Returning fullscreen changed the dock layout or main window geometry'
            assert [line.value() for line in selected.cursor_lines] == source_positions
            np.testing.assert_allclose(selected.getViewBox().viewRange()[0], zoomed, rtol=0, atol=1e-7)
            assert all(isValid(action) for action in workspace.navigation_actions.values()), 'Fullscreen exit deleted borrowed toolbar actions'
            assert not workspace._canvas_alignment_timer.signalsBlocked()
            return {'source_roles': ['machine', 'tdms'], 'selected_pane_only': True,
                    'pressure_A': float(selected.cursor_value(pressure.id, 0)),
                    'pressure_B': float(selected.cursor_value(pressure.id, 1)),
                    'cursor_A_source_index': index+1, 'zoom_ratio': .8,
                    'same_pane_curves_grid_and_docks_restored': True}
        recorder.check('REAL-SELECTED-GRAPH-FULLSCREEN-VALUES-RESTORE', selected_graph_fullscreen,
                       'graph_fullscreen', requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        def complete_workspace_fullscreen():
            from app_plot_workspace import PlotPane
            from shiboken6 import isValid
            window.tabs.setCurrentWidget(workspace)
            workspace.set_layout('Vertical split')
            workspace.set_time_lock(True)
            temperature = real_channel('tdms', '°C')
            pressure = real_channel('machine', 'Pa')
            speed = real_channel('machine', 'RPM')
            actuator = next(channel for channel in fixture_stores['machine'].channels.values()
                            if channel.attributes.get('actuator_binary'))
            first, second = workspace.visible_panes()
            first.bind([temperature.id, pressure.id])
            second.bind([speed.id, actuator.id])
            workspace.set_active(second)
            x, y = speed.values()
            finite = np.flatnonzero(np.isfinite(y))
            assert len(finite) > 3
            sample = int(finite[len(finite)//3])
            later = int(finite[2*len(finite)//3])
            assert sample+1 < later and np.isfinite(y[sample+1])
            for cursor, index in enumerate((sample, later)):
                second.lock_cursor(cursor, speed.id)
                second.cursor_locks[cursor]['index'] = index
            second.restore_ranges(second.state())
            events()
            panes = tuple(workspace.visible_panes())
            container = workspace.container
            slots = [workspace.grid.getItemPosition(workspace.grid.indexOf(pane)) for pane in panes]
            curves = [dict(pane.curves) for pane in panes]
            ids = [list(pane.ids) for pane in panes]
            ranges = [deepcopy(pane.getViewBox().viewRange()) for pane in panes]
            originals = {channel.id: (channel, channel.x, channel.y, channel.raw_y)
                         for channel in (temperature, pressure, speed, actuator)}
            docks, geometry = QtCore.QByteArray(window.saveState()), QtCore.QRect(window.geometry())
            controller = window.workspace_fullscreen
            assert controller.action.shortcut().toString() == 'F11'
            assert window.graph_fullscreen.action.shortcut().toString() == 'Shift+F11'
            checked_geometries = []
            try:
                # This is the same QAction as the graph-header control and F11.
                controller.action.trigger()
                for _ in range(4):events()
                host = controller.fullscreen_window
                assert controller.is_active and host.isVisible() and host.isFullScreen()
                assert controller.workspace is workspace
                assert controller.graph_scroll.widget() is container
                assert set(p for p in host.findChildren(PlotPane) if p.isVisible()) == set(panes)
                for index, pane in enumerate(panes):
                    assert pane.ids == ids[index] and pane.curves == curves[index]
                    np.testing.assert_allclose(pane.getViewBox().viewRange(), ranges[index], rtol=0, atol=1e-7)
                active = workspace.active
                stable_rects = [QtCore.QRect(pane.viewport().mapToGlobal(QtCore.QPoint()), pane.viewport().size())
                                for pane in panes]
                stable_ranges = [deepcopy(pane.getViewBox().viewRange()) for pane in panes]
                for index, pane in enumerate(panes):
                    bounds = pane.getViewBox().sceneBoundingRect()
                    point = bounds.topLeft()+QtCore.QPointF(bounds.width()*.4, bounds.height()*.45)
                    QtTest.QTest.mouseMove(pane.viewport(), pane.mapFromScene(point))
                    # Native/offscreen platforms differ in synthetic scene
                    # delivery. The real coalescer still resolves QCursor and
                    # widget ownership, including a queued stale leave event.
                    pane.scene().sigMouseMoved.emit(point)
                    for _ in range(3):events()
                    assert pane.hover_state is not None
                    assert 'Graph '+str(index+1) in controller.hover_readout.text()
                    assert len(controller.hover_readout.text().splitlines()) == 2
                    assert workspace.active is active, 'Read-only hover changed the active graph'
                    for number, graph in enumerate(panes):
                        assert QtCore.QRect(graph.viewport().mapToGlobal(QtCore.QPoint()), graph.viewport().size()) == stable_rects[number], 'Mouse hover moved a split graph canvas'
                        np.testing.assert_allclose(graph.getViewBox().viewRange(), stable_ranges[number], rtol=0, atol=1e-9)
                    QtTest.QTest.mouseMove(controller.exit_button, controller.exit_button.rect().center())
                    pane.scene().sigMouseMoved.emit(point)
                    for _ in range(3):events()
                    assert pane.hover_state is None and len(controller.hover_readout.text().splitlines()) == 1
                    for number, graph in enumerate(panes):
                        assert QtCore.QRect(graph.viewport().mapToGlobal(QtCore.QPoint()), graph.viewport().size()) == stable_rects[number]
                    checked_geometries.append([stable_rects[index].x(), stable_rects[index].y(),
                                               stable_rects[index].width(), stable_rects[index].height()])
                controller.pane_selector.setCurrentIndex(0)
                for _ in range(2):events()
                assert controller.pane is first
                assert {controller.readout.item(row, 0).data(QtCore.Qt.ItemDataRole.UserRole)
                        for row in range(controller.readout.rowCount())} == set(first.ids)
                controller.pane_selector.setCurrentIndex(1)
                for _ in range(2):events()
                assert controller.pane is second
                rows = {controller.readout.item(row, 0).data(QtCore.Qt.ItemDataRole.UserRole): row
                        for row in range(controller.readout.rowCount())}
                assert set(rows) == set(second.ids)
                assert controller.readout.item(rows[speed.id], 2).text() == f'{y[sample]:.6g}'
                workspace.navigation_actions['x'].trigger()
                before_x = np.array(second.getViewBox().viewRange()[0])
                before_y = [[list(view.viewRange()[1]) for view in pane.y_views()] for pane in panes]
                workspace.navigation_actions['in'].trigger()
                for _ in range(2):events()
                zoomed = second.getViewBox().viewRange()[0]
                assert np.isclose(np.ptp(zoomed)/np.ptp(before_x), .8)
                for number, pane in enumerate(panes):
                    np.testing.assert_allclose(pane.getViewBox().viewRange()[0], zoomed, rtol=0, atol=1e-7)
                    assert tuple(pane.getViewBox().state['mouseEnabled']) == (True, False)
                    for view, expected in zip(pane.y_views(), before_y[number]):
                        np.testing.assert_allclose(view.viewRange()[1], expected, rtol=0, atol=1e-7)
                second.step_cursor(0, 1)
                for _ in range(2):events()
                assert second.cursor_locks[0]['index'] == sample+1
                assert controller.readout.item(rows[speed.id], 2).text() == f'{y[sample+1]:.6g}'
                deliberate = [deepcopy(pane.state()) for pane in panes]
            finally:
                controller.exit()
                events()
            assert workspace.scroll.widget() is container
            assert tuple(workspace.visible_panes()) == panes
            assert window.saveState() == docks and window.geometry() == geometry
            for index, pane in enumerate(panes):
                assert workspace.grid.getItemPosition(workspace.grid.indexOf(pane)) == slots[index]
                assert pane.ids == ids[index] and pane.curves == curves[index]
                current = pane.state()
                for key in ('ids', 'bindings', 'cursors', 'cursor_locks', 'annotations'):
                    assert current.get(key) == deliberate[index].get(key), 'Fullscreen return changed '+key
                np.testing.assert_allclose(current['range'], deliberate[index]['range'], rtol=0, atol=1e-7)
            assert all(isValid(action) for action in workspace.navigation_actions.values())
            for channel_id, (channel, raw_x, raw_y, original_values) in originals.items():
                assert store.channels[channel_id] is channel and channel.x is raw_x and channel.y is raw_y
                assert channel.raw_y is original_values
            return {'source_roles': ['tdms', 'machine'], 'visible_panes': len(panes),
                    'units': ['°C', 'Pa', 'RPM', 'state'], 'hover_canvas_rectangles': checked_geometries,
                    'fixed_hover_geometry': True, 'same_container_panes_curves_and_grid_restored': True,
                    'cursor_A_source_index': sample+1, 'shared_X_zoom_ratio': .8,
                    'active_pane_values_and_raw_sources_preserved': True}
        recorder.check('REAL-WORKSPACE-FULLSCREEN-SPLITS-HOVER-RESTORE', complete_workspace_fullscreen,
                       'graph_fullscreen', requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        def two_signal_comparison_evidence():
            """Drive the actual one-click worker and inspect its complete outputs."""
            from unittest.mock import patch
            from openpyxl import load_workbook
            from matplotlib.backends.backend_pdf import PdfPages
            from matplotlib.text import Text
            from matplotlib.table import Table
            import re
            from app_dialog_history import remember_path

            reference = next(channel for channel in fixture_stores['machine'].channels.values()
                             if channel.unit == 'RPM' and 'feedback' in channel.name.casefold())
            candidate = next(channel for channel in fixture_stores['simulation'].channels.values()
                             if channel.unit == 'RPM' and channel.name == 'drum_speed')
            assert reference.dataset_id != candidate.dataset_id
            window.tabs.setCurrentWidget(workspace)
            window.portal.search.clear()
            window.portal.regex.setChecked(False)
            window.portal.rebuild()
            iterator = QtWidgets.QTreeWidgetItemIterator(window.portal.tree)
            chosen = {reference.id, candidate.id}
            while iterator.value():
                item = iterator.value()
                item.setSelected(item.data(0, QtCore.Qt.ItemDataRole.UserRole) in chosen)
                iterator += 1
            events()
            assert window.portal.selected() == [reference.id, candidate.id], 'The explicit cross-file selection was not retained'
            controls = window.signal_comparison
            before_views = deepcopy([view.state() for view in window.workspaces])
            before_channel_ids, before_dataset_ids = set(store.channels), set(store.datasets)
            originals = {channel.id: (channel, channel.x, channel.y, channel.raw_y,
                sha256_file(samples/next(entry['filename'] for entry in manifest['files']
                    if entry['role'] == role)),
                hashlib.sha256(channel.x.tobytes()+channel.y.tobytes()).hexdigest())
                for role, channel in (('machine', reference), ('simulation', candidate))}
            metadata = deepcopy(window.engineering.session['metadata'])
            test_note = 'Acceptance evidence: compare supported measurements; no test limits evaluated.'
            window.engineering.session['metadata']['Observations'] = test_note
            destination = output/'two-signal-report'
            destination.mkdir(exist_ok=True)
            assert remember_path('SignalComparisonReport', destination/'requested.pdf')
            before_files = set(destination.iterdir())
            warnings = []
            rendered_pages = []
            save_pdf_page = PdfPages.savefig
            def observe_saved_pdf_page(pdf, figure, *args, **kwargs):
                # Observe the actual text artists supplied to the shipped PDF
                # backend and still call that backend. No PDF reader library
                # or alternate test renderer is required in frozen packages.
                texts = [artist.get_text() for artist in figure.findobj(match=Text)]
                for table in figure.findobj(match=Table):
                    texts.extend(cell.get_text().get_text() for cell in table.get_celld().values())
                value = save_pdf_page(pdf, figure, *args, **kwargs)
                rendered_pages.append('\n'.join(texts))
                return value
            started = time.monotonic()
            try:
                # Only a failure-message boundary is made nonblocking. The
                # selected QAction, engine, background worker, report writers
                # and results dialog all run as shipped in this distribution.
                with patch.object(QtWidgets.QMessageBox, 'warning',
                        side_effect=lambda parent, title, message, *args: warnings.append(str(message))), \
                        patch.object(PdfPages, 'savefig', observe_saved_pdf_page):
                    controls.action.trigger()
                    assert window.busy and window.workers, 'Compare did not start its background operation'
                    wait_for_acceptance_workers(window, lambda: acceptance_worker_events(app), timeout=360)
                    events()
                assert warnings == [], 'The real comparison reported errors: '+repr(warnings)
                result, bundle = controls.last_result, controls.last_bundle
                assert result and result['status'] == 'compared' and bundle
                assert result['reference']['channel_id'] == reference.id
                assert result['candidate']['channel_id'] == candidate.id
                assert result['alignment']['mode'] == 'elapsed'
                assert result['metrics']['unit'] == 'RPM'
                independent = verify_comparison_interval_metrics(result)
                paths = {key: Path(bundle[key]) for key in ('pdf', 'workbook', 'evidence')}
                assert all(path.parent.resolve() == destination.resolve() and path.is_file() and path.stat().st_size > 0
                           for path in paths.values())
                assert set(destination.iterdir())-before_files == set(paths.values()), 'A report transaction left partial or unowned files'
                evidence = json.loads(paths['evidence'].read_text(encoding='utf-8'))
                exported = evidence['signal_comparison']
                assert exported == result, 'The JSON companion differs from the displayed comparison'
                assert evidence['report']['Metadata']['Observations'] == test_note
                assert evidence['raw_interval_sample_count'] == len(reference.x)+len(candidate.x)
                assert len(evidence['raw_interval_samples']) == evidence['raw_interval_sample_count']
                assert len(exported['differences']) == independent['complete_intervals']
                # The workbook keeps native numbers, not six-digit display
                # strings, and retains every sample, interval and event.
                book = load_workbook(paths['workbook'], read_only=True, data_only=True)
                try:
                    required = {'Active interval', 'Cursor summary', 'Pair comparison',
                                'Difference intervals', 'Activity events', 'Sources'}
                    assert required <= set(book.sheetnames)
                    assert book['Active interval'].max_row-1 == evidence['raw_interval_sample_count']
                    assert book['Difference intervals'].max_row-1 == len(result['differences'])
                    assert book['Activity events'].max_row-1 == len(result['events'])
                    rows = iter(book['Pair comparison'].values)
                    headers = next(rows)
                    assert 'Field' in headers and 'Value' in headers
                    fi, vi = headers.index('Field'), headers.index('Value')
                    metric_cells = {row[fi]: row[vi] for row in rows}
                    for key in ('compared_duration_min', 'coverage_fraction', 'mean_bias',
                                'mean_absolute_difference', 'rmse', 'max_absolute_difference'):
                        value = metric_cells['metrics.'+key]
                        assert isinstance(value, (int, float)) and not isinstance(value, bool)
                        assert np.isclose(value, result['metrics'][key], rtol=1e-12, atol=1e-10), 'Workbook rounded '+key
                finally:
                    book.close()
                pdf_bytes = paths['pdf'].read_bytes()
                assert pdf_bytes.startswith(b'%PDF-') and b'%%EOF' in pdf_bytes[-1024:]
                page_count = len(re.findall(rb'/Type\s*/Page\b', pdf_bytes))
                assert page_count == len(rendered_pages), 'The saved PDF page objects differ from actual backend writes'
                pdf_text = '\n'.join(rendered_pages)
                pdf_words = ' '.join(pdf_text.split())
                assert reference.name in pdf_words and candidate.name in pdf_words
                assert 'elapsed' in pdf_text.casefold() and 'coverage' in pdf_text.casefold()
                assert 'Acceptance evidence' in pdf_words and 'no test limits evaluated' in pdf_words
                for label in ('Time-weighted average difference', 'Root mean square difference',
                              'Largest absolute difference', 'Supported comparison coverage'):
                    assert label in pdf_words, 'The PDF omitted the paired result '+label
                assert evidence['pdf_pages'] == page_count and page_count >= 4
                assert controls.result_dialog is not None and controls.result_dialog.isVisible()
                assert reference.name in controls.result_dialog.sources.text()
                assert candidate.name in controls.result_dialog.sources.text()
                shown = {controls.result_dialog.metrics.item(row, 0).text():
                         controls.result_dialog.metrics.item(row, 1).text()
                         for row in range(controls.result_dialog.metrics.rowCount())}
                assert shown['Root mean square difference'] == f'{result["metrics"]["rmse"]:.6g} RPM'
                assert shown['Supported comparison coverage'] == f'{100*result["metrics"]["coverage_fraction"]:.6g}%'
                button_texts = {button.text() for button in controls.result_dialog.findChildren(QtWidgets.QPushButton)}
                assert {'Open Report', 'Open Workbook', 'Open Folder'} <= button_texts
                assert window.last_export == paths['pdf']
                assert controls.action.isEnabled()
                assert [view.state() for view in window.workspaces] == before_views, 'Pair comparison changed the original plotted workspace'
                assert set(store.channels) == before_channel_ids and set(store.datasets) == before_dataset_ids
                for role, channel in (('machine', reference), ('simulation', candidate)):
                    original, raw_x, raw_y, raw_original, file_hash, array_hash = originals[channel.id]
                    assert store.channels[channel.id] is original and channel.x is raw_x and channel.y is raw_y
                    assert channel.raw_y is raw_original
                    assert hashlib.sha256(channel.x.tobytes()+channel.y.tobytes()).hexdigest() == array_hash
                    filename = next(entry['filename'] for entry in manifest['files'] if entry['role'] == role)
                    assert sha256_file(samples/filename) == file_hash
                return {'source_roles': ['machine', 'simulation'],
                    'reference': reference.name, 'candidate': candidate.name, 'alignment': 'elapsed',
                    'unit': 'RPM', 'raw_samples': evidence['raw_interval_sample_count'],
                    'complete_difference_intervals': len(result['differences']),
                    'complete_activity_events': len(result['events']), 'pdf_pages': page_count,
                    'independent_interval_integrals': independent,
                    'bundle': {key: {'filename': path.name, 'bytes': path.stat().st_size,
                                    'sha256': sha256_file(path)} for key, path in paths.items()},
                    'real_QAction_worker_and_results_dialog': True,
                    'raw_sources_and_original_workspaces_unchanged': True,
                    'seconds_with_output_inspection': round(time.monotonic()-started, 3)}
            finally:
                window.engineering.session['metadata'] = metadata
                if controls.result_dialog is not None:
                    controls.result_dialog.close()
                events()
        recorder.check('REAL-TWO-SIGNAL-COMPARISON-EVIDENCE', two_signal_comparison_evidence,
                       'comparison_report', requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        def combinations_and_model_removal():
            assert len(fixture_stores) == 4, 'Import must complete before file combinations'
            rows = []
            roles = tuple(fixture_stores)
            for count in range(1, len(roles)+1):
                for group in combinations(roles, count):
                    mixed = ChannelStore()
                    for role in group:
                        mixed.datasets.update(fixture_stores[role].datasets)
                        mixed.channels.update(fixture_stores[role].channels)
                    expected = sum(len(fixture_stores[role].channels) for role in group)
                    assert len(mixed.channels) == expected
                    workspace.set_layout('Single')
                    for pane in workspace.panes:
                        pane.bind([])
                    representatives = [next(iter(fixture_stores[role].channels)) for role in group]
                    workspace.active.bind(representatives)
                    events()
                    assert set(workspace.active.ids) == set(representatives)
                    first = store.channels[representatives[0]]
                    workspace.active.lock_cursor(0, first.id)
                    workspace.active.step_cursor(0, 1)
                    assert workspace.active.measurements(), 'Combination cursor inspection did not produce readouts'
                    workspace.set_visible(representatives, False)
                    assert not any(pane.ids for pane in workspace.visible_panes())
                    workspace.set_visible(representatives, True)
                    events()
                    plotted = {channel_id for pane in workspace.visible_panes() for channel_id in pane.ids}
                    assert plotted == set(representatives), 'Combination deselect/reselect crossed file partitions'
                    for role in group:
                        isolated = ChannelStore()
                        isolated.datasets.update(mixed.datasets)
                        isolated.channels.update(mixed.channels)
                        removed = isolated.remove_datasets(fixture_stores[role].datasets)
                        target = set(fixture_stores[role].channels)
                        assert removed == target, 'Removal crossed a dataset partition'
                        assert set(isolated.channels) == set(mixed.channels)-target
                        for channel_id, channel in isolated.channels.items():
                            assert channel is mixed.channels[channel_id], 'Unrelated channel identity changed'
                    rows.append({'roles': list(group), 'channels': expected, 'isolated_removals': count,
                                 'representative_plot_channels': len(representatives),
                                 'cursor_step_and_deselect_reselect': True})
            assert len(rows) == 15
            return rows
        recorder.check('ALL-15-FILE-COMBINATIONS', combinations_and_model_removal, 'real_recordings',
                       requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        def file_selection_and_removal():
            assert len(fixture_stores) == 4
            window.tabs.setCurrentIndex(0)
            workspace.set_layout('Single')
            for pane in workspace.panes:
                pane.bind([])
            window.sync_visibility()
            window.portal.search.clear()
            window.portal.rebuild()
            tree = window.portal.tree
            assert tree.columnCount() == 2, 'Data Portal must contain only Channel and Unit columns'
            headers = [tree.headerItem().text(index) for index in range(tree.columnCount())]
            assert headers == ['Channel', 'Unit'], 'Unexpected Data Portal column or header'
            iterator = QtWidgets.QTreeWidgetItemIterator(tree)
            while iterator.value():
                item = iterator.value()
                assert all(tree.itemWidget(item, column) is None for column in range(tree.columnCount())), 'Inline file-action widgets consume the channel columns'
                iterator += 1
            observed = []
            for role, other in fixture_stores.items():
                ids = set(other.channels)
                dataset_id = next(iter(other.datasets))
                root = window.portal.dataset_roots[dataset_id]
                root.setCheckState(0, QtCore.Qt.CheckState.Checked)
                events()
                plotted = {channel_id for pane in workspace.visible_panes() for channel_id in pane.ids}
                assert plotted == ids, f'{role} Select All plotted other files or missed its own channels'
                root = window.portal.dataset_roots[dataset_id]
                root.setCheckState(0, QtCore.Qt.CheckState.Unchecked)
                events()
                assert not any(pane.ids for pane in workspace.visible_panes()), 'Deselect All left orphan traces'
                observed.append({'role': role, 'selected_channels': len(ids)})
            first, second = list(fixture_stores.values())[:2]
            selected = next(iter(first.channels))
            workspace.set_layout('Single')
            workspace.active.bind([selected])
            preserved = store.channels[selected]
            window.drop_recordings(list(second.datasets))
            events()
            assert store.channels[selected] is preserved and selected in workspace.active.ids
            assert not set(second.channels).intersection(store.channels)
            # Re-add the exact untouched source objects for subsequent checks.
            store.datasets.update(second.datasets)
            with store.batch():
                for channel in second.channels.values():
                    store.add(channel)
            return {'files': observed, 'portal_columns': headers, 'inline_item_widgets': 0}
        recorder.check('PER-FILE-CHECKBOX-AND-CLOSE-DATASET', file_selection_and_removal, 'dataset_isolation',
                       requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        def tdms_template_lifecycle():
            assert 'tdms' in fixture_stores
            other = fixture_stores['tdms']
            path = output/'templates'/'tdms-parameters.json'
            templates = ParameterTemplates(path)
            unknown = templates.unknown(other)
            assert unknown, 'The first TDMS import must require review'
            review = ParameterReview(unknown, window)
            try:
                entries = review.entries()
            finally:
                review.close()
            templates.save(entries)
            restarted = ParameterTemplates(path)
            assert not restarted.unknown(other), 'A saved TDMS template was not reused'
            probe = next(iter(other.channels.values()))
            from dataclasses import replace
            new = replace(probe, id='acceptance-unknown-tdms', name='New unknown parameter',
                attributes={**probe.attributes, 'tdms_template_key': 'acceptance/new/unknown'})
            mixed = ChannelStore()
            mixed.channels.update(other.channels)
            mixed.add(new)
            assert [channel.id for channel in restarted.unknown(mixed)] == [new.id]
            assert identity(new) not in restarted.entries
            restarted.apply(mixed)
            for channel_id in other.channels:
                np.testing.assert_array_equal(mixed.channels[channel_id].x, other.channels[channel_id].x)
                # Templates retain original readings even for actuator views.
                original = other.channels[channel_id].raw_y
                if original is not None:
                    np.testing.assert_array_equal(mixed.channels[channel_id].raw_y, original)
            return {'reviewed_parameters': len(entries), 'unknown_only_review': 1}
        recorder.check('TDMS-FIRST-REVIEW-REUSE-UNKNOWN-ONLY', tdms_template_lifecycle, 'templates',
                       requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))

        synthetic = []
        def synthetic_sources():
            ds = Dataset('Acceptance edge cases', pd.DataFrame({'time_min': [0., 1., 1., 2., 4.],
                 'reading': [10., 20., 30., np.nan, 50.], 'literal': ['0', '01', '', 'NaN', '004']}))
            store.datasets[ds.id] = ds
            definitions = [
                Channel('Duplicate readings', [0, 1, 1, 2, 4], [10, 20, 30, np.nan, 50],
                    unit='°C', coordinate_unit='min', dataset_id=ds.id),
                Channel('Discrete valve', [0, 1, 1, 2, 4], [0, 1, 0, np.nan, 1],
                    unit='state', kind='discrete', coordinate_unit='min', dataset_id=ds.id),
                Channel('Masked pressure', [0, 1, 1, 2, 4], [100, 200, 250, 300, 500],
                    unit='Pa', coordinate_unit='min', dataset_id=ds.id, masks=[(2, 2)]),
                Channel('Verified power', np.arange(121)/60, np.full(121, 1000.),
                    unit='W', coordinate_unit='min', dataset_id=ds.id, attributes={'match_percent': 93}),
                Channel('Unknown physical unit', np.arange(121)/60, np.ones(121),
                    coordinate_unit='min', dataset_id=ds.id),
            ]
            with store.batch():
                for channel in definitions:
                    synthetic.append(store.add(channel))
            return {'channels': len(synthetic), 'cases': ['duplicates', 'missing', 'masks', 'discrete', 'unknown units']}
        recorder.check('SYNTHETIC-SOURCE-EDGE-CASES', synthetic_sources, 'synthetic')

        def locked_identity():
            assert len(synthetic) == 5
            window.tabs.setCurrentIndex(0)
            workspace.set_layout('Single')
            pane = workspace.active
            observed = []
            for channel in synthetic[:3]:
                pane.bind([channel.id])
                pane.lock_cursor(0, channel.id)
                # Start before the duplicates, then step through both original
                # source rows; assigning a timestamp alone cannot identify them.
                pane.cursor_locks[0]['index'] = 0
                pane.restore_ranges({**pane.state(), 'cursor_locks': [dict(pane.cursor_locks[0]), None]})
                for expected_index in (1, 2, 3, 4):
                    pane.step_cursor(0, 1)
                    events()
                    lock = pane.cursor_locks[0]
                    assert lock['index'] == expected_index, 'Arrow stepping skipped a duplicate source observation'
                    expected = channel.values()[1][expected_index]
                    measured = pane.measurements()[channel.id]['A']
                    assert np.isclose(measured, expected, equal_nan=True), 'Locked readout disagrees with its exact sample'
                    marker_x, marker_y = pane.cursor_markers[0].getData()
                    if np.isfinite(expected):
                        assert len(marker_y) == 1 and np.isclose(marker_y[0], expected), 'Locked marker uses another duplicate row'
                        assert np.isclose(marker_x[0], channel.values()[0][expected_index])
                    else:
                        assert (not pane.cursor_markers[0].isVisible() or not np.isfinite(marker_y).any()), 'A masked or missing sample has a finite visible marker'
                    _, summary = interval_tables(store, workspace.state())
                    assert np.isclose(summary.iloc[0]['A'], expected, equal_nan=True), 'Export endpoint changed sample identity'
                    saved = pane.state()
                    pane.restore(saved)
                    assert pane.cursor_locks[0]['index'] == expected_index
                    assert np.isclose(pane.measurements()[channel.id]['A'], expected, equal_nan=True)
                    window._readouts()
                    if np.isfinite(expected):
                        assert np.isclose(float(window.readout.item(0, 2).text()), expected), 'HUD disagrees with its sample marker'
                    else:
                        assert window.readout.item(0, 2).text() == 'Unknown'
                # Crossing and clamping must retain nonnegative span and valid
                # source indices; no source rows are removed to obtain this.
                pane.cursor_lines[1].setValue(0)
                pane.cursor_lines[0].setValue(1e9)
                events()
                stats = pane.measurements()[channel.id]
                assert stats['ΔX'] >= 0 and 0 <= pane.cursor_locks[0]['index'] < len(channel.x)
                observed.append({'channel': channel.name, 'source_samples': len(channel.x), 'stepped_rows': 4})
            return observed
        recorder.check('LOCKED-SAMPLE-MARKER-HUD-EXPORT-RESTORE', locked_identity, 'cursor_identity',
                       requires=('SYNTHETIC-SOURCE-EDGE-CASES',))

        def header_and_context():
            workspace.set_layout('Vertical split')
            for pane, channel in zip(workspace.visible_panes(), synthetic):
                pane.bind([channel.id])
            events()
            bar = workspace.navigation_bar
            assert bar.parentWidget() is not window, 'Navigation is still mounted in the global toolbar'
            button = bar.widgetForAction(workspace.navigation_actions['in'])
            assert button is not None
            distance = abs(button.mapToGlobal(button.rect().center()).y()-window.tabs.tabBar().mapToGlobal(window.tabs.tabBar().rect().center()).y())
            assert distance <= 16, f'Navigation is {distance} pixels away from the graph tab header'
            bars = [toolbar for toolbar in window.findChildren(QtWidgets.QToolBar)
                    if toolbar.objectName() in ('graphNavigation', 'AnalysisNavigation', 'ComparisonNavigation') and toolbar.isVisible()]
            assert len(bars) == 1, 'Duplicate primary graph navigation bars are visible'
            pane = workspace.visible_panes()[1]
            menu = pane.create_context_menu()
            try:
                labels = [action.text().replace('&', '').lower() for action in menu.actions() if not action.isSeparator()]
                text = '\n'.join(labels)
                for term in ('pan', 'cursor', 'zoom xy', 'zoom x', 'zoom y', 'snap to cycles', 'revert zoom', 'comments', 'copy', 'export'):
                    assert term in text, 'Missing pane context task: '+term
                assert workspace.active is pane, 'Context menu acts on a different pane'
                assert not pane.getPlotItem().ctrl.fftCheck.isEnabled(), 'Unsupported native FFT remains user-accessible'
                # Even a stale programmatic toggle must not raise a Qt callback.
                pane.getPlotItem().ctrl.fftCheck.setChecked(True)
                pane.getPlotItem().ctrl.fftCheck.setChecked(False)
                events()
            finally:
                menu.close()
            return {'header_center_difference': distance, 'visible_navigation_bars': len(bars), 'context_actions': labels}
        recorder.check('GRAPH-HEADER-CONTEXT-SAFE-NATIVE-FFT', header_and_context, 'navigation',
                       requires=('SYNTHETIC-SOURCE-EDGE-CASES',))

        def assert_ranges(expected=None):
            panes = [pane for pane in workspace.visible_panes() if pane.coordinate == ('time', 'min') and pane.ids]
            assert panes
            ranges = [np.asarray(pane.getViewBox().viewRange()[0]) for pane in panes]
            expected = ranges[0] if expected is None else np.asarray(expected)
            for bounds in ranges:
                np.testing.assert_allclose(bounds, expected, rtol=0, atol=1e-7,
                    err_msg='Lock Time Axes numeric endpoints diverged')
            return [bounds.tolist() for bounds in ranges]

        def layout_theme_geometry(layout, theme, size):
            window.resize(*size)
            window.theme.apply(theme)
            workspace.set_layout(layout)
            _, columns = LAYOUTS[layout]
            real_temperature = next((channel for channel in fixture_stores.get('tdms', ChannelStore()).channels.values()
                if channel.coordinate_kind == 'time' and channel.coordinate_unit == 'min' and channel.unit == '°C'), synthetic[0])
            real_pressure = next((channel for channel in fixture_stores.get('machine', ChannelStore()).channels.values()
                if channel.coordinate_kind == 'time' and channel.coordinate_unit == 'min' and channel.unit == 'Pa'), synthetic[2])
            for index, pane in enumerate(workspace.visible_panes()):
                row, column = divmod(index, columns)
                # In stacked/grid layouts every column contains both narrow
                # single-unit and wide independent multi-unit axis gutters.
                ids = [real_temperature.id] if (row+column) % 2 == 0 else [real_pressure.id, synthetic[3].id, synthetic[1].id]
                pane.bind(ids)

            def settle_painted_axes():
                previous = None
                stable = 0
                for attempt in range(12):
                    app.processEvents()
                    # Offscreen scroll children may not paint with the window.
                    # Grab each pane so tick labels discover their real widths.
                    for pane in workspace.visible_panes():
                        pane.grab()
                    QtTest.QTest.qWait(15)
                    app.processEvents()
                    geometry = tuple(round(viewport_time_x(pane, value), 5)
                        for pane in workspace.visible_panes()
                        for value in pane.getViewBox().viewRange()[0])
                    stable = stable+1 if geometry == previous else 0
                    previous = geometry
                    if stable >= 3:
                        return attempt+1
                raise AssertionError('Painted axis widths and queued canvas alignment did not settle')

            workspace.active = workspace.panes[0]
            workspace.set_time_lock(False)
            workspace.active.setXRange(.5, 1.5, padding=0)
            workspace.set_time_lock(True)
            paint_passes = {'time_lock': settle_painted_axes()}
            adopted = assert_ranges([.5, 1.5])
            pixel_checks = {'time_lock': time_pixel_alignment(workspace.visible_panes(), columns)}
            workspace.visible_panes()[-1].setXRange(.25, 1.75, padding=0)
            paint_passes['pan'] = settle_painted_axes()
            assert_ranges([.25, 1.75])
            pixel_checks['pan'] = time_pixel_alignment(workspace.visible_panes(), columns)
            for mode in ('x', 'y', 'xy', 'pan'):
                workspace.navigation_actions[mode].trigger()
                assert workspace.navigation_mode == mode
                assert workspace.navigation_actions[mode].isChecked()
                for pane in workspace.visible_panes():
                    enabled = list(pane.getViewBox().state['mouseEnabled'])
                    assert enabled == [mode != 'y', mode != 'x'], 'Menu/header mode and gesture mode differ'
            workspace.navigate('x')
            workspace.navigate('in')
            paint_passes['zoom_in'] = settle_painted_axes()
            assert_ranges()
            pixel_checks['zoom_in'] = time_pixel_alignment(workspace.visible_panes(), columns)
            workspace.navigate('out')
            workspace.navigate('reset')
            paint_passes['reset'] = settle_painted_axes()
            assert_ranges()
            pixel_checks['reset'] = time_pixel_alignment(workspace.visible_panes(), columns)
            saved = workspace.state()
            workspace.restore(saved)
            paint_passes['restore'] = settle_painted_axes()
            bounds = assert_ranges()
            pixel_checks['restore'] = time_pixel_alignment(workspace.visible_panes(), columns)
            heights = [pane.getViewBox().sceneBoundingRect().height() for pane in workspace.visible_panes()]
            assert min(heights) >= 150, f'{layout} at {size} crushed drawing heights to {heights}'
            assert len(workspace.visible_panes()) == LAYOUTS[layout][0]*LAYOUTS[layout][1]
            return {'layout': layout, 'theme': theme, 'window_size': size,
                    'adopted_ranges': adopted, 'restored_ranges': bounds, 'drawing_heights': heights,
                    'paint_settle_passes': paint_passes, 'time_pixel_alignment': pixel_checks,
                    'scroll_maximum': workspace.scroll.verticalScrollBar().maximum()}
        for layout in LAYOUTS:
            for theme in ('Dark', 'Light'):
                for size in ((1024, 680), (1536, 960)):
                    recorder.check(f'LAYOUT-{layout}-{theme}-{size[0]}',
                        lambda layout=layout, theme=theme, size=size: layout_theme_geometry(layout, theme, size),
                        'layout_time_modes', requires=('SYNTHETIC-SOURCE-EDGE-CASES',))

        def domain_independence():
            workspace.set_layout('Horizontal split')
            pane, other = workspace.visible_panes()
            pane.bind([synthetic[3].id])
            frequency = store.add(transform(synthetic[3], 'FFT'))
            other.bind([frequency.id])
            workspace.active = pane
            workspace.set_time_lock(True)
            other.setXRange(.1, .4, padding=0)
            pane.setXRange(.5, 1.5, padding=0)
            events()
            np.testing.assert_allclose(other.getViewBox().viewRange()[0], [.1, .4], rtol=0, atol=1e-7)
            other.bind([synthetic[2].id])
            workspace.set_time_lock(False)
            pane.setXRange(.1, .6, padding=0)
            other.setXRange(1.1, 1.8, padding=0)
            events()
            assert not np.allclose(pane.getViewBox().viewRange()[0], other.getViewBox().viewRange()[0])
            return {'frequency_independent': True, 'unlocked_time_independent': True}
        recorder.check('NON-TEMPORAL-AND-UNLOCKED-INDEPENDENCE', domain_independence, 'time_lock',
                       requires=('SYNTHETIC-SOURCE-EDGE-CASES',))

        def numerical_units_and_guards():
            power = synthetic[3]
            integral = transform(power, 'Cumulative integration')
            assert np.isclose(integral.y[-1], 120000.), 'Minute-clock integration did not use seconds'
            derived = transform(power, 'Differentiation')
            np.testing.assert_allclose(derived.y, 0, rtol=0, atol=1e-8)
            for channel in synthetic[:3]:
                try:
                    strict_clock(channel)
                except ValueError as exc:
                    recorder.result['expected_safety_rejections'].append({'scenario': 'strict-clock', 'source': channel.name, 'message': str(exc)})
                else:
                    raise AssertionError('Irregular/duplicate/masked clock silently accepted')
            for operation in ('Butterworth low-pass', 'FFT', 'Welch PSD'):
                try:
                    transform(synthetic[0], operation, cutoff=.1)
                except ValueError as exc:
                    recorder.result['expected_safety_rejections'].append({'scenario': operation, 'message': str(exc)})
                else:
                    raise AssertionError('Changing a cutoff bypassed an irregular-clock rejection')
            try:
                formula(store, f'Ch("{synthetic[0].id}") + Ch("{power.id}")')
            except ValueError as exc:
                recorder.result['expected_safety_rejections'].append({'scenario': 'mixed-clock-formula', 'message': str(exc)})
            else:
                raise AssertionError('Mixed clocks combined without explicit resampling')
            return {'integral_watt_seconds': float(integral.y[-1]), 'derivative_max_abs': float(np.max(np.abs(derived.y)))}
        recorder.check('NUMERICAL-UNITS-AND-SAFETY-REJECTIONS', numerical_units_and_guards, 'processing',
                       requires=('SYNTHETIC-SOURCE-EDGE-CASES',))

        def processing_preflight_controls():
            def select(channel):
                window.portal.search.clear()
                window.portal.rebuild()
                iterator = QtWidgets.QTreeWidgetItemIterator(window.portal.tree)
                while iterator.value():
                    item = iterator.value()
                    item.setSelected(item.data(0, QtCore.Qt.ItemDataRole.UserRole) == channel.id)
                    iterator += 1
                window.update_processing_preflight(reset_default=True)
                events()
            window.operation.setCurrentText('Butterworth low-pass')
            select(synthetic[3])
            cutoff = window.parameter_fields['cutoff'].value()
            assert 0 < cutoff < .5, 'The default cutoff is outside the verified Nyquist range'
            assert window.apply_button.isEnabled()
            assert 'nyquist' in window.parameter_help.text().lower()
            result = transform(synthetic[3], 'Butterworth low-pass', cutoff=cutoff)
            assert len(result.x) == len(synthetic[3].x)
            select(synthetic[0])
            assert not window.apply_button.isEnabled(), 'Irregular-clock Apply remains enabled'
            assert 'resampl' in window.parameter_help.text().lower(), 'Missing actionable resampling explanation'
            window.operation.setCurrentText('Cumulative integration')
            select(synthetic[4])
            assert 'unknown physical' in window.parameter_help.text().lower()
            unknown = transform(synthetic[4], 'Cumulative integration')
            assert 'wh' not in unknown.unit.lower(), 'Unknown source units were presented as known energy'
            return {'regular_cutoff_hz': cutoff, 'irregular_apply_blocked': True,
                    'unknown_unit_explanation': window.parameter_help.text()}
        recorder.check('PROCESSING-PREFLIGHT-VALID-DEFAULT-AND-EXPLANATIONS', processing_preflight_controls,
                       'processing', requires=('SYNTHETIC-SOURCE-EDGE-CASES',))

        def synthetic_gestures():
            window.tabs.setCurrentIndex(0)
            window.resize(1536, 960)
            workspace.set_layout('Single')
            pane = workspace.active
            pane.bind([synthetic[3].id, synthetic[2].id])
            pane.set_cursor_mode('off')
            rows = []
            def ranges():
                return (np.asarray(pane.getViewBox().viewRange()[0]),
                        [np.asarray(view.viewRange()[1]) for view in pane.y_views()])
            for mode in ('x', 'y', 'xy'):
                for gesture in ('wheel', 'box'):
                    workspace.navigate('reset')
                    workspace.navigate(mode)
                    events()
                    before_x, before_y = ranges()
                    bounds = pane.getViewBox().sceneBoundingRect()
                    if gesture == 'wheel':
                        position = pane.mapFromScene(bounds.center())
                        event = QtGui.QWheelEvent(QtCore.QPointF(position), QtCore.QPointF(pane.viewport().mapToGlobal(position)),
                            QtCore.QPoint(), QtCore.QPoint(0, 120), QtCore.Qt.MouseButton.NoButton,
                            QtCore.Qt.KeyboardModifier.NoModifier, QtCore.Qt.ScrollPhase.NoScrollPhase, False)
                        app.sendEvent(pane.viewport(), event)
                    else:
                        start = pane.mapFromScene(bounds.topLeft()+QtCore.QPointF(bounds.width()*.25, bounds.height()*.25))
                        end = pane.mapFromScene(bounds.topLeft()+QtCore.QPointF(bounds.width()*.75, bounds.height()*.75))
                        QtTest.QTest.mouseMove(pane.viewport(), start)
                        # QTest's mouseMove ``delay`` advances its synthetic
                        # timestamps on some Qt platforms without waiting in
                        # real time. Keep normal scene rate limiting enabled
                        # and pace a real drag so its moves reach GraphicsScene.
                        QtTest.QTest.qWait(25)
                        QtTest.QTest.mousePress(pane.viewport(), QtCore.Qt.MouseButton.LeftButton, pos=start)
                        QtTest.QTest.qWait(25)
                        QtTest.QTest.mouseMove(pane.viewport(), start+(end-start)/2)
                        QtTest.QTest.qWait(25)
                        QtTest.QTest.mouseMove(pane.viewport(), end)
                        QtTest.QTest.qWait(25)
                        QtTest.QTest.mouseRelease(pane.viewport(), QtCore.Qt.MouseButton.LeftButton, pos=end)
                    events()
                    after_x, after_y = ranges()
                    if mode == 'y':
                        np.testing.assert_allclose(after_x, before_x, rtol=0, atol=1e-7)
                    else:
                        assert np.ptp(after_x) < np.ptp(before_x)*.95, f'{mode} {gesture}: X gesture did not zoom the selected axis'
                    for first, last in zip(before_y, after_y):
                        if mode == 'x':
                            np.testing.assert_allclose(last, first, rtol=0, atol=1e-7)
                        else:
                            assert np.ptp(last) < np.ptp(first)*.95, 'Y gesture did not zoom an independent axis'
                    rows.append({'mode': mode, 'gesture': gesture})
            QtTest.QTest.keyClick(pane, QtCore.Qt.Key.Key_Escape)
            events()
            assert workspace.navigation_mode == 'pan' and workspace.navigation_actions['pan'].isChecked()
            assert list(pane.getViewBox().state['mouseEnabled']) == [True, True]
            pane.set_cursor_mode('dual')
            return {'qt_synthetic_gestures': rows, 'escape_shared_mode': 'pan',
                    'physical_mouse_and_dpi_checks': 'Unrun; see manual checks'}
        recorder.check('QT-WHEEL-BOX-KEYBOARD-SHARED-MODES', synthetic_gestures, 'navigation',
                       requires=('SYNTHETIC-SOURCE-EDGE-CASES',))

        def exact_workspace():
            assert len(fixture_stores) == 4
            expected_source_ids = {channel_id for other in fixture_stores.values() for channel_id in other.channels}
            assert expected_source_ids <= set(store.channels), 'A source partition is missing before preservation check'
            window.resize(1536, 960)
            workspace.set_layout('Single')
            pane = workspace.active
            pane.bind([synthetic[0].id])
            pane.lock_cursor(0, synthetic[0].id)
            pane.cursor_locks[0]['index'] = 1
            pane.restore_ranges(pane.state())
            views = [view.state() for view in window.workspaces]
            path = save_workspace(output/'source-preserving-workspace.json', store, views, window.ui_state())
            restored, saved_views, _ = load_workspace(path)
            assert set(restored.channels) == set(store.channels)
            assert set(restored.datasets) == set(store.datasets)
            for channel_id, channel in store.channels.items():
                other = restored.channels[channel_id]
                np.testing.assert_array_equal(other.x, channel.x)
                np.testing.assert_array_equal(other.y, channel.y)
                if channel.raw_y is not None:
                    np.testing.assert_array_equal(other.raw_y, channel.raw_y)
                np.testing.assert_array_equal(other.masks, channel.masks)
                assert other.shift == channel.shift
                assert other.unit == channel.unit and other.coordinate_unit == channel.coordinate_unit
            for dataset_id, dataset in store.datasets.items():
                pd.testing.assert_frame_equal(restored.datasets[dataset_id].frame, dataset.frame,
                                             check_exact=True, check_dtype=True)
                assert restored.datasets[dataset_id].frame.attrs == dataset.frame.attrs
            assert saved_views == views
            size = path.stat().st_size+sum(child.stat().st_size for child in path.parent.glob(path.stem+'.*.npz'))
            # Large exact-preserving snapshots are diagnostic outputs only. Keep
            # bounded reports/screenshots in CI artifacts, not 200+ MB per run.
            companion = path.parent/json.loads(path.read_text(encoding='utf-8'))['arrays']
            path.unlink()
            companion.unlink()
            return {'channels': len(restored.channels), 'datasets': len(restored.datasets),
                    'exact_source_cells': True, 'snapshot_bytes': size, 'snapshot_removed_after_verification': True}
        recorder.check('EXACT-REAL-SOURCE-WORKSPACE-ROUNDTRIP', exact_workspace, 'persistence',
                       requires=('REAL-FIXTURE-INTEGRITY-IMPORT', 'PER-FILE-CHECKBOX-AND-CLOSE-DATASET'))

        def exports_agree():
            assert fixture_stores
            workspace.set_layout('Single')
            representatives = [next(iter(other.channels)) for other in fixture_stores.values()]
            pane = workspace.active
            # Compatible time channels from different runs may intentionally be
            # overlaid; no alignment or matching-condition inference is made.
            compatible = [channel_id for channel_id in representatives if store.channels[channel_id].coordinate_kind == 'time']
            pane.bind(compatible+[synthetic[3].id])
            pane.cursor_lines[0].setValue(1)
            pane.cursor_lines[1].setValue(1.5)
            events()
            state = workspace.state()
            data, summary = interval_tables(store, state)
            for suffix in ('csv', 'xlsx', 'pdf', 'png'):
                path = output/('active-interval.'+suffix)
                exported, exported_summary = export_report(path, store, state)
                assert path.is_file() and path.stat().st_size > 100
                pd.testing.assert_frame_equal(exported, data)
                pd.testing.assert_frame_equal(exported_summary, summary)
            csv = pd.read_csv(output/'active-interval.csv')
            workbook = pd.read_excel(output/'active-interval.xlsx', sheet_name=None)
            assert list(workbook) == ['Active interval', 'Cursor summary']
            assert len(csv) == len(data) == len(workbook['Active interval'])
            disk_summary = workbook['Cursor summary']
            for field in ('cursor_A_X', 'cursor_B_X', 'A', 'B', 'ΔY'):
                np.testing.assert_allclose(disk_summary[field], summary[field], equal_nan=True)
            assert set(disk_summary.channel_id) == set(summary.channel_id)
            from openpyxl import load_workbook
            book = load_workbook(output/'active-interval.xlsx')
            sheet = book['Cursor summary']
            headers = [cell.value for cell in sheet[1]]
            color_column = headers.index('match_percent')+1
            assert any(sheet.cell(row, color_column).fill.fill_type == 'solid' for row in range(2, sheet.max_row+1))
            return {'summary_channels': len(summary), 'source_rows': len(data),
                    'formats': ['csv', 'xlsx', 'pdf', 'png'], 'cursor_and_match_metadata_verified': True}
        recorder.check('REAL-INTERVAL-EXPORT-FORMATS-AND-AGREEMENT', exports_agree, 'exports',
                       requires=('REAL-FIXTURE-INTEGRITY-IMPORT', 'PER-FILE-CHECKBOX-AND-CLOSE-DATASET',
                                 'SYNTHETIC-SOURCE-EDGE-CASES'))

        def sources_unchanged():
            assert manifest
            for entry in manifest['files']:
                assert sha256_file(samples/entry['filename']) == entry['sha256'], 'Original source file changed'
            return {'originals_verified': len(manifest['files'])}
        recorder.check('SOURCE-FILES-UNCHANGED', sources_unchanged, 'real_recordings',
                       requires=('REAL-FIXTURE-INTEGRITY-IMPORT',))
        window.grab().save(str(output/'accepted-workspace.png'))
    except Exception:
        recorder.result['runner_errors'].append({'id': recorder.current, 'error': traceback.format_exc()})
        recorder.save()
    finally:
        if window is not None:
            window.close()
            events()
            QtCore.QCoreApplication.sendPostedEvents(None, QtCore.QEvent.Type.DeferredDelete)
        sys.excepthook = old_hook
