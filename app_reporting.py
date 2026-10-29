"""Direct exports of the visible workspace and its active cursor intervals."""
from pathlib import Path
import numpy as np
import pandas as pd
from app_presentation import match_color
from app_signal_processing import IntervalSummary
from app_time_units import presentation_statistics
from app_cursor_identity import interval_rows, interval_statistics, locked_index
from app_operation_context import checkpoint
from app_channel_identity import channel_label


def literal_workbook_text(book):
    """Report strings are literal labels/notes, never spreadsheet formulas.

    Measurements and statistics are calculated before export. Preserve text
    starting with '=' or Excel error tokens without changing numeric precision.
    """
    for sheet in book:
        checkpoint('Retaining literal report text in '+sheet.title)
        for row in sheet.iter_rows():
            if row and row[0].row % 8192 == 0:checkpoint()
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'


def match_scorecard(score):
    """Express verified relative differences as a bounded similarity percentage.

    Match = max(0, 100 - abs(relative difference percent)). Missing, unverified
    and zero-denominator metrics retain no percentage and receive no accent.
    The engineering engine's signed differences and original values are retained.
    """
    result=score.drop(columns=['status'],errors='ignore').copy()
    difference=pd.to_numeric(score.get('percent',pd.Series(index=score.index,dtype=float)).astype(str).str.replace('%','',regex=False),errors='coerce')
    verified=score.get('status',pd.Series('',index=score.index)).isin(['GREEN','YELLOW','RED'])
    result['match_percent']=(100-difference.abs()).clip(lower=0,upper=100).where(verified & np.isfinite(difference))
    return result


def interval_tables(store,workspace):
    from app_plot_workspace import LAYOUTS
    rows,cols=LAYOUTS[workspace['layout']];frames=[];summary=[]
    for pane_number,view in enumerate(workspace['panes'][:rows*cols],1):
        checkpoint('Collecting active interval samples',pane_number-1,rows*cols)
        a,b=view.get('cursors',[0,0]);low,high=sorted((a,b))
        for channel_id in view['ids']:
            c=store.channels[channel_id];x,y=c.values();locks=view.get('cursor_locks',[]);mask=interval_rows(c,a,b,locks)
            stats=interval_statistics(c,IntervalSummary(c),a,b,locks,view.get('snap',True));va,vb=stats['A'],stats['B']
            metadata={'pane':pane_number,'channel_id':c.id,'channel':channel_label(store,c),'original_channel':c.name,'unit':c.unit,'coordinate_unit':c.coordinate_unit,
                      'cursor_A_X':a,'cursor_B_X':b,**presentation_statistics(c,stats)}
            if workspace.get('statistics_scope')=='Whole recording' and len(x):
                whole=IntervalSummary(c).query(float(x[0]),float(x[-1]))
                metadata.update({key:whole[key] for key in ('Min','Max','Mean')});metadata['min_max_mean_scope']='Whole recording'
            for index,lock in enumerate(view.get('cursor_locks',[])):
                if lock and lock.get('channel') in store.channels:
                    metadata[f'cursor_{("A","B")[index]}_locked_channel']=channel_label(store,lock['channel'])
                    metadata[f'cursor_{("A","B")[index]}_locked_channel_id']=lock['channel']
                    metadata[f'cursor_{("A","B")[index]}_sample_index']=locked_index(store.channels[lock['channel']],lock)
            metadata['ΔY']=vb-va
            metadata['match_percent']=c.attributes.get('match_percent',np.nan) if match_color(c.attributes.get('match_percent')) else np.nan
            if c.attributes.get('actuator_binary'):
                from app_actuators import state_text
                metadata.update(value_interpretation='actuator',A_state=state_text(va),B_state=state_text(vb))
                scope=IntervalSummary(c).query(float(x[0]),float(x[-1])) if workspace.get('statistics_scope')=='Whole recording' and len(x) else IntervalSummary(c).query(a,b)
                metadata['ON_time']=scope['Mean']*scope['Coverage']*scope['ΔX'];metadata['ON_fraction']=scope['Mean'];metadata['known_time']=scope['Coverage']*scope['ΔX']
            summary.append(metadata)
            frames.append(pd.DataFrame({'pane':pane_number,'channel_id':c.id,'channel':channel_label(store,c),'original_channel':c.name,'source_sample_index':np.arange(len(x))[mask],'coordinate':x[mask],
                                        'value':y[mask],'unit':c.unit,'coordinate_unit':c.coordinate_unit,
                                        'cursor_A_X':a,'cursor_B_X':b,'match_percent':metadata['match_percent'],**({'original_value':c.raw_y[mask]} if c.raw_y is not None else {})}))
    if not summary:raise ValueError('Check at least one channel before exporting')
    return pd.concat(frames,ignore_index=True),pd.DataFrame(summary)


def export_report(path,store,workspace,scorecard=None,*,title='Signal analysis report',metadata=None):
    """One save dialog, no configuration wizard; never interpolate exported samples."""
    from app_workspace_io import render_report
    from app_operation_context import checkpoint
    from app_atomic_export import atomic_export
    checkpoint('Collecting active cursor interval');path=Path(path);data,summary=interval_tables(store,workspace)
    if path.suffix.lower() not in ('.csv','.xlsx','.pdf','.png'):raise ValueError('Choose PDF, Excel, CSV or PNG')
    comparison=(workspace.get('report_kind')=='comparison' or workspace.get('workspace_name')=='Comparison')
    if comparison and path.suffix.lower()=='.pdf':
        from app_comparison_report import export_comparison_bundle
        export_comparison_bundle(path,store,workspace,metadata,scorecard,title=title,_prepared=(data,summary))
        return data,summary
    if comparison and path.suffix.lower()=='.xlsx':
        import json
        from app_comparison_report import build_evidence,write_workbook,clean
        from app_engineering_reports import staged_report_bundle
        evidence,tables=build_evidence(store,workspace,data,summary,metadata,scorecard,title)
        companion=path.with_suffix('.evidence.json')
        with staged_report_bundle([path,companion]) as staged:
            write_workbook(staged[path],data,summary,evidence,tables,scorecard)
            checkpoint('Writing complete comparison JSON evidence')
            staged[companion].write_text(json.dumps(clean(evidence),indent=2,ensure_ascii=False,allow_nan=False)+'\n',encoding='utf-8')
            checkpoint('Committing comparison workbook and evidence')
        return data,summary
    with atomic_export(path) as temporary:
        if path.suffix.lower()=='.csv':data.to_csv(temporary,index=False)
        elif path.suffix.lower()=='.xlsx':
            from openpyxl.styles import PatternFill
            with pd.ExcelWriter(temporary,engine='openpyxl') as writer:
                checkpoint('Writing interval samples');data.to_excel(writer,sheet_name='Active interval',index=False)
                checkpoint('Writing cursor summary');summary.to_excel(writer,sheet_name='Cursor summary',index=False)
                if scorecard is not None:scorecard.to_excel(writer,sheet_name='Comparison',index=False)
                for sheet in writer.book:
                    checkpoint('Formatting '+sheet.title);sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
                    headers=[cell.value for cell in sheet[1]]
                    if 'match_percent' in headers:
                        column=headers.index('match_percent')+1
                        for row in range(2,sheet.max_row+1):
                            if row%8192==0:checkpoint()
                            cell=sheet.cell(row,column);color=match_color(cell.value)
                            if color:cell.fill=PatternFill('solid',fgColor=color.lstrip('#'));cell.number_format='0.0"%"'
                literal_workbook_text(writer.book)
        else:render_report(temporary,store,workspace,title=title,scorecard=scorecard)
    return data,summary
